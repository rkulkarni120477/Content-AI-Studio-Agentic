"""Excel export of the prompt list (GET /prompt-library/prompts/export.xlsx).

Replaces the console's client-side CSV serializer: the backend now builds the
workbook with the exact filters the list page sends, so the download always
matches what the user is looking at. Column contract mirrors the retired CSV
export, with ``category`` renamed to ``workflow`` (the list page's Workflow
filter; "Category" now means cluster in the UI).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

EXPECTED_COLUMNS = (
    "id", "parent_id", "parent_title", "title", "prompt", "description",
    "workflow", "visibility", "teams", "tags", "variable_names", "created_by",
    "created_at", "updated_at", "last_used_at", "review_count", "review_avg",
    "version_count",
)


def _create_pipeline_prompt(client, headers, name="xlsx_probe", **overrides):
    payload = {
        "name": name,
        "description": "xlsx export characterization",
        "component_type": "cdd",
        "system_prompt": "You are a course designer.",
        "user_prompt_template": "Design {{topic}}.",
    }
    payload.update(overrides)
    resp = client.post("/api/v1/prompts", json=payload, headers=headers)
    assert resp.status_code == 201, resp.text
    return resp.json()


def _export_rows(client, headers, **params):
    resp = client.get("/api/v1/prompt-library/prompts/export.xlsx",
                      params=params, headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="prompts-export-' in resp.headers["content-disposition"]
    ws = load_workbook(BytesIO(resp.content), read_only=True).active
    return list(ws.iter_rows(values_only=True))


class TestExportXlsx:
    def test_header_row_matches_contract(self, client, auth_headers):
        rows = _export_rows(client, auth_headers, kind="pipeline", roots_only="1")
        assert rows[0] == EXPECTED_COLUMNS

    def test_pipeline_row_carries_workflow_label_and_content(self, client, auth_headers):
        created = _create_pipeline_prompt(client, auth_headers)
        rows = _export_rows(client, auth_headers, kind="pipeline", roots_only="1")
        by_id = {r[0]: r for r in rows[1:]}
        row = by_id.get(created["id"])
        assert row is not None, "created prompt missing from export"
        cols = dict(zip(EXPECTED_COLUMNS, row, strict=True))
        assert cols["workflow"] == "CDD"
        assert "Design {{topic}}." in (cols["prompt"] or "")
        assert cols["created_by"] == "test_admin"
        assert cols["version_count"] >= 1

    def test_cas_category_filter_applies(self, client, auth_headers):
        created = _create_pipeline_prompt(client, auth_headers, name="xlsx_filter_probe")
        included = _export_rows(client, auth_headers, kind="pipeline",
                                roots_only="1", cas_category="CDD")
        excluded = _export_rows(client, auth_headers, kind="pipeline",
                                roots_only="1", cas_category="Blueprint")
        assert any(r[0] == created["id"] for r in included[1:])
        assert not any(r[0] == created["id"] for r in excluded[1:])

    def test_empty_result_still_returns_valid_workbook(self, client, auth_headers):
        rows = _export_rows(client, auth_headers, kind="pipeline",
                            roots_only="1", q="no_prompt_matches_this_9f2c")
        assert rows[0] == EXPECTED_COLUMNS
        assert len(rows) == 1
