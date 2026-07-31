"""DLU (worksheet-based) CDD detection, splitting, and multi-sheet XLSX export.

Covers:
- is_dlu_cdd() / split_cdd_worksheets() in promptops_app/parsers/cdd_parser.py
- build_xlsx_worksheets() + _safe_sheet_name() in promptops_app/exporters/xlsx_exporter.py

The DLU XLSX path emits one sheet per worksheet (Overview + Worksheet 1..N),
with markdown tables rendered as real cells. Standard CDDs are not DLU and use
the existing single-sheet export unchanged.
"""

from __future__ import annotations

from openpyxl import load_workbook

from promptops_app.parsers.cdd_parser import is_dlu_cdd, split_cdd_worksheets
from promptops_app.exporters.xlsx_exporter import build_xlsx_worksheets, _safe_sheet_name

DLU_CS = """# BLOCK 2 BLUEPRINT
## Sample Block

**Project:** Test Project

## WORKSHEET 1: BLOCK OVERVIEW
Overview body line one.

## WORKSHEET 2: INSTRUCTIONAL SEQUENCE MAP
| Day | Day Title | ACS Code |
|-----|-----------|----------|
| 1 | Intro | AM.I.E.K1 |
| 2 | Next | AM.I.E.K2 |

## WORKSHEET 3: ACS CODE REGISTRY
Some prose describing coverage."""

STANDARD_CS = """## Course Details
A standard CDD.

## Course Structure
Module 1: Intro
Module 2: Advanced"""

# ── Shared format fixtures ────────────────────────────────────────────────────
# The worksheet label format is set by whichever prompt generated the CDD, so
# detection must not depend on markdown heading syntax. These mirror the cases
# in frontend/src/utils/__tests__/cddWorksheets.test.js — keep the two in sync.

# Real shape produced by the "Block Blueprint" prompt from 2026-07-30 (CDDs
# 74-77): bold labels, zero markdown headings. Used to render as a flat blob.
BOLD_CS = """**Block Blueprint for Block 2: Aircraft Drawings**

---

**Worksheet 1: Block Overview**

- **Block Title:** Aircraft Drawings
- **Domain:** General Studies

---

**Worksheet 2: Source File Inventory**

- **Course Syllabus/Calendar:** Available
- **ACS Codes:** Available

---

**Worksheet 3: ACS Code Registry**

- **AM.I.E.S6:** Skill-type ACS code"""

PLAIN_CS = """Block Overview Document

Worksheet 1 — Block Overview
Body one.

Worksheet 2 . Source File Inventory
Body two."""

# A standard CDD that merely mentions a worksheet in prose must not flip to DLU.
PROSE_MENTION_CS = """## Course Structure
Module 1: Intro. Learners complete Worksheet 1 before the lab session, as
recorded in Worksheet 1 of the prior block.
Module 2: Advanced"""

# Table of contents printed before the real worksheets — naive splitting would
# emit near-empty worksheets at the TOC lines.
TOC_CS = """# Block Blueprint

- **Worksheet 1: Block Overview**
- **Worksheet 2: Source File Inventory**

---

**Worksheet 1: Block Overview**

The real overview body, which is substantially longer than the TOC entry.

**Worksheet 2: Source File Inventory**

The real inventory body, also substantially longer than its TOC entry."""


def test_is_dlu_cdd_detection():
    assert is_dlu_cdd(DLU_CS) is True
    assert is_dlu_cdd(STANDARD_CS) is False
    assert is_dlu_cdd("") is False


def test_is_dlu_cdd_accepts_non_heading_label_formats():
    """Bold and plain labels are DLU too — the prompt decides the format."""
    assert is_dlu_cdd(BOLD_CS) is True
    assert is_dlu_cdd(PLAIN_CS) is True


def test_is_dlu_cdd_ignores_prose_mentions():
    """One mid-sentence mention must not reshape a standard CDD."""
    assert is_dlu_cdd(PROSE_MENTION_CS) is False
    # A single bold label on its own is still not enough evidence.
    assert is_dlu_cdd("Intro text\n\n**Worksheet 1: Only One**\n\nBody.") is False
    # ...but a single markdown heading is (the original, unchanged rule).
    assert is_dlu_cdd("Intro text\n\n## WORKSHEET 1: ONLY ONE\n\nBody.") is True


def test_split_cdd_worksheets_bold_labels():
    sections = split_cdd_worksheets(BOLD_CS)
    labels = [lbl for lbl, _ in sections]
    assert labels[0] == "Overview"
    assert len(sections) == 4  # Overview + 3 worksheets
    assert labels[1].startswith("1 ")
    assert any("ACS" in lbl for lbl in labels)
    # The model's own bold label line is preserved verbatim in the body.
    assert sections[1][1].startswith("**Worksheet 1: Block Overview**")
    # Content is split at the right boundaries.
    assert "Aircraft Drawings" in sections[1][1]
    assert "Course Syllabus/Calendar" in sections[2][1]
    assert "AM.I.E.S6" in sections[3][1]


def test_split_cdd_worksheets_prefers_real_section_over_toc():
    sections = split_cdd_worksheets(TOC_CS)
    # Two worksheets, not four — the TOC lines are discarded.
    assert len([lbl for lbl, _ in sections if lbl != "Overview"]) == 2
    bodies = {lbl: content for lbl, content in sections}
    assert any("real overview body" in c for c in bodies.values())
    assert any("real inventory body" in c for c in bodies.values())


def test_split_cdd_worksheets_overview_first_then_worksheets():
    sections = split_cdd_worksheets(DLU_CS)
    labels = [lbl for lbl, _ in sections]
    assert labels[0] == "Overview"
    assert len(sections) == 4  # Overview + 3 worksheets
    # Short, acronym-preserving labels.
    assert labels[1].startswith("1 ")
    assert any("ACS" in lbl for lbl in labels)  # not "Acs"
    # Each worksheet section keeps its own heading (boundary marker).
    for lbl, content in sections[1:]:
        assert content.lstrip().upper().startswith("## WORKSHEET") or "WORKSHEET" in content.upper()


def test_split_cdd_worksheets_empty_for_standard():
    assert split_cdd_worksheets(STANDARD_CS) == []


def test_build_xlsx_worksheets_one_sheet_per_section_with_table_cells():
    sections = split_cdd_worksheets(DLU_CS)
    buf = build_xlsx_worksheets("Sample Block", sections)
    wb = load_workbook(buf)

    # One sheet per section, Overview first.
    assert len(wb.sheetnames) == 4
    assert wb.sheetnames[0] == "Overview"

    # Sheet names are Excel-safe.
    for name in wb.sheetnames:
        assert len(name) <= 31
        assert not any(ch in name for ch in '\\/?*[]:')
        assert not (name.startswith("'") or name.endswith("'"))

    # Worksheet 2's markdown table became real cells (header row + 2 day rows).
    seq_sheet = next(wb[n] for n in wb.sheetnames if n.startswith("2 "))
    rows = [
        [c for c in row if c not in (None, "")]
        for row in seq_sheet.iter_rows(values_only=True)
    ]
    header = next(r for r in rows if r and r[0] == "Day")
    assert "Day Title" in header and "ACS Code" in header
    day_values = {r[0] for r in rows if len(r) >= 3 and str(r[0]) in ("1", "2")}
    assert day_values == {"1", "2"}


def test_build_xlsx_worksheets_empty_falls_back():
    # No sections -> single-sheet fallback, still a valid workbook.
    buf = build_xlsx_worksheets("Empty", [])
    wb = load_workbook(buf)
    assert len(wb.sheetnames) >= 1


def test_safe_sheet_name_rules():
    used = set()
    # Illegal chars stripped, length capped, word-boundary truncation.
    long_name = _safe_sheet_name("3 ACS Code Registry With Orphan Check", used)
    used.add(long_name)
    assert len(long_name) <= 31
    assert "ACS" in long_name

    # Leading/trailing apostrophes removed (Excel rejects them).
    apos = _safe_sheet_name("'Weird' Name'", used)
    assert not apos.startswith("'")
    assert not apos.endswith("'")

    # Illegal characters replaced.
    illegal = _safe_sheet_name("A/B:C*D", set())
    assert not any(ch in illegal for ch in '\\/?*[]:')

    # Duplicate names are disambiguated.
    u = set()
    a = _safe_sheet_name("Overview", u); u.add(a)
    b = _safe_sheet_name("Overview", u)
    assert a != b
