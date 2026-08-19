"""Alignment of the generated block-wide workbook with the AIM reference sample.

Guards the shape a reviewer actually diffs against
``data/Block_02_Block_Blueprint_AK.xlsx``. Column identity and ORDER both matter —
the reference is pasted into an existing review workflow, so a reordered column is
as disruptive as a missing one.

Deliberate deviations from the reference, asserted here so they stay deliberate:
  * "Notes" is ours (the REDUCE narrative) with no reference counterpart → last.
  * "AIM SME Comments" is reviewer-fill and blank by design (0/20 in the reference).
  * "Targets for Quick Check" / "Summative Exam Item Cluster" are derived in code and
    state the data gap where the reference's AKTR / exam-blueprint sources are not
    ingested — they must never contain invented percentages or item ranges.
"""
from __future__ import annotations

from promptops_app.services.block_wide_generator import BlockWideGenerator
from promptops_app.services.block_wide_generator import EXTENSION_MISSING
from promptops_app.services.block_wide_service import (
    _DAY_TABLE_HEADER, _day_table_from_rows, _self_review_lines,
)


# The reference workbook's "4_Day-by-Day Map" header, verbatim, in order.
REFERENCE_COLUMNS = [
    "Day #", "Topic Label (exact from calendar)", "Handbook Reference (calendar, exact)",
    "Handbook Edition as Cited", "ACS Codes Today", "Concept Type",
    "Concept Type Explanation", "Concept Scope", "Learn-While-Doing", "How it is Applied",
    "Hangar Activity", "Projects  Today", "Assessment Today", "Targets for Quick Check",
    "Summative Exam Item Cluster", "Source Files for This Day", "Proposed Learning Objective",
    "Known Misconceptions / Student Difficulties", "Interactive Candidate",
    "Potential Storyline Interactive Type", "Interactive Content",
    "Storyline Source Asset Status", "Interactive Scope", "Interactive Rationale",
    "Job Aid Candidate", "Job Aid Content Type", "Job Aid Content Description",
    "Job Aid Source Reference", "Academian Questions", "AIM SME Comments",
]

#: reference label → our (shorter) label, where they differ by wording only.
_ALIAS = {
    "Day #": "Day",
    "Topic Label (exact from calendar)": "Topic",
    "Handbook Reference (calendar, exact)": "Handbook Reference",
    "Handbook Edition as Cited": "Handbook Edition",
    "ACS Codes Today": "ACS",
    "How it is Applied": "How It Is Applied",
    "Projects  Today": "Projects Today",
    "Source Files for This Day": "Source Files",
    "Proposed Learning Objective": "Learning Objective",
    "Known Misconceptions / Student Difficulties": "Misconceptions",
    "Potential Storyline Interactive Type": "Interactive Type",
    "Job Aid Content Type": "Job Aid Type",
    "Job Aid Content Description": "Job Aid Description",
}


def _expected():
    return [_ALIAS.get(c, c) for c in REFERENCE_COLUMNS]


def test_every_reference_column_is_present_in_reference_order():
    """The whole point: a reviewer diffing our sheet against the reference should
    find the same columns in the same order."""
    expected = _expected()
    assert _DAY_TABLE_HEADER[:len(expected)] == expected, (
        "day-table columns diverged from the AIM reference:\n"
        f"  expected: {expected}\n  actual:   {_DAY_TABLE_HEADER[:len(expected)]}"
    )


def test_only_notes_is_added_beyond_the_reference():
    """Extra columns push everything after them out of alignment, so any addition
    must sit past the reference's last column."""
    assert _DAY_TABLE_HEADER[len(_expected()):] == ["Notes"]


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #
def _reduce(digests, acs_registry=None):
    enumerate_summary = {
        "total_days": 2, "client_id": "aim", "block": "Block 2",
        "declared_acs": ["AM.I.B.K1", "AM.I.B.K2"],
        "days": [
            {"day_number": 1, "topic": "Aircraft drawings", "acs_codes": ["AM.I.B.K1"],
             "projects_today": [], "assessment_today": [], "handbook_reference": "FAA-H-8083-30B",
             "handbook_edition": "FAA-H-8083-30B", "source_files_today": ["B2D1.pdf"],
             "hangar_activity_today": []},
            {"day_number": 2, "topic": "Methods", "acs_codes": ["AM.I.B.K2"],
             "projects_today": ["Project 2-1"], "assessment_today": ["Quiz 1"],
             "handbook_reference": "", "handbook_edition": "", "source_files_today": [],
             "hangar_activity_today": []},
        ],
    }
    return BlockWideGenerator(llm=lambda m, s, u: "{}").reduce(
        enumerate_summary, digests, deliverable="cdd", tier="draft",
        acs_registry=acs_registry,
    )


def _day_cells(result):
    rows = next(s for s in result.sections if s.get("key") == "day_table")["rows"]
    lines = _day_table_from_rows(rows)
    header = [c.strip() for c in lines[0].strip("|").split("|")]
    out = {}
    for line in lines[2:]:                      # [1] is the |---| separator
        cells = [c.strip() for c in line.strip("|").split("|")]
        assert len(cells) == len(header), "row/header cell count mismatch"
        # The Day # cell renders as "Day 5" (matching the AIM reference), but these
        # tests index by the bare number — normalise here so the day format stays a
        # rendering detail rather than something every assertion has to know.
        out[cells[0].removeprefix("Day ").strip()] = dict(zip(header, cells))
    return out


_OK_DIGEST = {
    "day_number": 1, "concept_type": "Conceptual",
    "concept_type_explanation": "First-encounter knowledge acquisition.",
    "concept_scope": "Title-block contents, universal numbering system.",
    "derived_objective": "Interpret title blocks.", "misconceptions": ["Views are photos"],
    "digest_status": "ok", "storyline_source_asset_status": "AVAILABLE",
}
_REGISTRY = [
    {"acs_code": "AM.I.B.K1", "high_miss": "79.6% (rank #1)", "priority": "APPLY/ANALYZE"},
    {"acs_code": "AM.I.B.K2", "high_miss": "NO AKTR DATA", "priority": ""},
]


def test_every_row_has_one_cell_per_column():
    """A cell-count drift silently shifts every value one column left or right,
    which reads as plausible data in the wrong field."""
    cells = _day_cells(_reduce([_OK_DIGEST, {"day_number": 2, "digest_status": "failed"}], _REGISTRY))
    # Keys are normalised by _day_cells; the rendered "Day N" format is pinned in
    # test_block_wide_coverage.test_day_table_from_rows_produces_one_line_per_day.
    assert set(cells) == {"1", "2"}


def test_concept_scope_carries_the_digest_field():
    cells = _day_cells(_reduce([_OK_DIGEST], _REGISTRY))
    assert cells["1"]["Concept Scope"] == "Title-block contents, universal numbering system."


def test_concept_scope_flags_a_failed_day_instead_of_going_blank():
    cells = _day_cells(_reduce([{"day_number": 1, "digest_status": "failed"}], _REGISTRY))
    assert "REVIEW NEEDED" in cells["1"]["Concept Scope"]


def test_quick_check_targets_report_registry_facts_and_never_invent_them():
    cells = _day_cells(_reduce([_OK_DIGEST, {"day_number": 2, "digest_status": "ok"}], _REGISTRY))
    day1 = cells["1"]["Targets for Quick Check"]
    assert "79.6% (rank #1)" in day1 and "APPLY/ANALYZE" in day1
    # Day 2's code has no AKTR figure — must say so, not imply it was assessed low.
    assert "NO AKTR DATA" in cells["2"]["Targets for Quick Check"]


def test_quick_check_targets_handle_a_day_with_no_acs_codes():
    result = _reduce([_OK_DIGEST], _REGISTRY)
    rows = next(s for s in result.sections if s.get("key") == "day_table")["rows"]
    rows[0]["acs_codes"] = []
    BlockWideGenerator._apply_assessment_columns(rows, _REGISTRY)
    assert "No ACS codes" in rows[0]["quick_check_targets"]


def test_quick_check_targets_survive_an_absent_registry():
    """The registry is a separate DIS worksheet; a partial bundle must not crash the
    render or fabricate priorities."""
    cells = _day_cells(_reduce([_OK_DIGEST], acs_registry=None))
    assert "NO AKTR DATA" in cells["1"]["Targets for Quick Check"]


def test_summative_cluster_reports_the_gap_rather_than_estimating():
    """The reference's item ranges come from an exam blueprint that is not ingested;
    inventing "Items ~1-5" would be indistinguishable from a real range."""
    cells = _day_cells(_reduce([_OK_DIGEST], _REGISTRY))
    value = cells["1"]["Summative Exam Item Cluster"]
    assert "REVIEW NEEDED" in value
    assert "~" not in value, "must not fabricate an item range"


def test_aim_sme_comments_is_blank_by_design():
    cells = _day_cells(_reduce([_OK_DIGEST], _REGISTRY))
    assert cells["1"]["AIM SME Comments"] == ""


# --------------------------------------------------------------------------- #
# XLSX export — key/value worksheets
# --------------------------------------------------------------------------- #
def test_kv_bullets_export_as_label_value_cells():
    """Worksheets 1 and 5 are key-value grids in the reference. Written verbatim the
    markdown bullet landed as one string in column A ("- **Block:** Block 2"), which
    is unusable as a spreadsheet."""
    import openpyxl
    from promptops_app.exporters.xlsx_exporter import build_xlsx_worksheets

    md = ("## WORKSHEET 1: BLOCK OVERVIEW\n\n"
          "- **Block:** Block 2\n- **Total Days:** 20\n- **Grading Policy:** \n"
          "- plain bullet with no label\n")
    ws = openpyxl.load_workbook(
        build_xlsx_worksheets("Block 2", [("1 Block Overview", md)]), data_only=True
    )["1 Block Overview"]
    pairs = {r[0]: r[1] for r in ws.iter_rows(min_row=1, max_col=2, values_only=True) if r[0]}

    assert pairs["Block"] == "Block 2"
    assert pairs["Total Days"] == "20"
    assert pairs["Grading Policy"] in (None, "")      # empty value keeps its label row
    # A bullet with no bold label still renders as content, without its "- " marker.
    assert "plain bullet with no label" in pairs
    # And no cell contains raw markdown.
    assert not any("**" in str(k) for k in pairs)


def test_markdown_tables_still_export_as_grids():
    """The day table must keep exporting as real cells — the new bullet branch runs
    only for non-table lines."""
    import openpyxl
    from promptops_app.exporters.xlsx_exporter import build_xlsx_worksheets

    md = "## WORKSHEET 4\n\n| Day | Topic |\n|---|---|\n| 1 | Drawings |\n"
    ws = openpyxl.load_workbook(
        build_xlsx_worksheets("t", [("4 Day-by-day Map", md)]), data_only=True
    )["4 Day-by-day Map"]
    grid = [r for r in ws.iter_rows(min_row=1, max_col=2, values_only=True) if r[0]]
    assert ("Day", "Topic") in grid and ("1", "Drawings") in grid


# --------------------------------------------------------------------------- #
# Five-pass self-review (AIM's Blueprint spec ends every document with it)
# --------------------------------------------------------------------------- #
def _clean_coverage() -> dict:
    return {"total_days": 20, "days_in_output": 20, "missing_days": [],
            "declared_acs": ["AM.I.B.K1", "AM.I.E.K6"],
            "covered_acs": ["AM.I.B.K1", "AM.I.E.K6"],
            "orphan_acs": [], "failed_days": [], "thin_days": []}


def test_self_review_emits_exactly_five_numbered_passes():
    body = [l for l in _self_review_lines(_clean_coverage(), 5) if l.startswith(tuple("12345"))]
    assert [l.split(".")[0] for l in body] == ["1", "2", "3", "4", "5"]


def test_a_clean_run_raises_no_warning_marker():
    assert "⚠" not in "\n".join(_self_review_lines(_clean_coverage(), 5))


def test_each_pass_reports_the_defect_it_is_responsible_for():
    cov = _clean_coverage() | {"days_in_output": 19, "missing_days": [13],
                               "covered_acs": ["AM.I.B.K1"], "orphan_acs": ["AM.I.E.K6"],
                               "thin_days": [7]}
    out = _self_review_lines(cov, 4)
    assert "4/5 worksheets" in out[2]
    assert "19/20 day rows" in out[3] and "day 13" in out[3]
    assert "1/2 declared ACS" in out[4] and "AM.I.E.K6" in out[4]
    assert "day 7" in out[5]
    assert out.count("") == 1                      # one blank separator, no stray lines


def test_day_lists_are_prose_not_python_reprs():
    """f-string interpolation of the coverage lists yields "[13]" — the form that
    reads as a bug in a document a client reviews."""
    cov = _clean_coverage() | {"missing_days": [13, 14], "thin_days": [7]}
    text = "\n".join(_self_review_lines(cov, 5))
    assert "[" not in text and "]" not in text
    assert "days 13, 14" in text and "day 7" in text


def test_the_specificity_pass_never_claims_a_check_it_did_not_run():
    """The one pass that cannot be computed says so. A green line nobody earned is
    worse than an honest gap, because it is the line a reviewer would have trusted."""
    line = _self_review_lines(_clean_coverage(), 5)[6]
    assert line.startswith("5. **Specificity**")
    assert "NOT machine-checkable" in line and "reviewer" in line


def test_self_review_survives_a_coverage_report_missing_every_optional_key():
    """Renders from whatever the coverage dict actually has — a KeyError here would
    cost the document its whole review section."""
    out = _self_review_lines({}, 0)
    assert len(out) == 7 and "0/5 worksheets" in out[2]


# --------------------------------------------------------------------------- #
# Prompt-declared additional columns are ADDITIVE — the safety property
# --------------------------------------------------------------------------- #
def _one_row() -> list[dict]:
    return [{"day_number": 1, "topic": "Drawings", "acs_codes": ["AM.I.B.K1"],
             "extensions": {"Stage": "Interpret"}}]


def test_a_declared_column_appends_and_never_displaces_the_reference_columns():
    """The whole reason this mechanism is affordable AND safe: a prompt can add a
    column, never rename, reorder or blank one the renderer computes."""
    base = _day_table_from_rows(_one_row())
    with_extra = _day_table_from_rows(_one_row(), ["Stage"])
    base_header = base[0].strip("|").split("|")
    extra_header = with_extra[0].strip("|").split("|")
    assert extra_header[:len(_DAY_TABLE_HEADER)] == base_header
    assert extra_header[len(_DAY_TABLE_HEADER):] == [" Stage "]
    assert with_extra[2].strip("|").split("|")[:len(_DAY_TABLE_HEADER)] == \
        base[2].strip("|").split("|")


def test_the_separator_row_widens_with_the_header():
    """A separator that does not match the header count makes the whole table stop
    rendering as a table in a strict renderer."""
    out = _day_table_from_rows(_one_row(), ["Stage", "Pilot"])
    assert out[1].count("---") == len(_DAY_TABLE_HEADER) + 2


def test_every_row_still_has_one_cell_per_column_with_extras():
    out = _day_table_from_rows(_one_row(), ["Stage", "Pilot"])
    width = len(_DAY_TABLE_HEADER) + 2
    for line in out[2:]:
        assert len(line.strip("|").split("|")) == width


def test_a_declared_column_with_no_value_states_the_gap():
    out = _day_table_from_rows([{"day_number": 1, "topic": "x"}], ["Stage"])
    assert "REVIEW NEEDED — not returned" in out[2]


def test_no_declaration_is_byte_identical_to_before_the_feature():
    assert _day_table_from_rows(_one_row()) == _day_table_from_rows(_one_row(), [])
    assert _day_table_from_rows(_one_row()) == _day_table_from_rows(_one_row(), None)


def test_a_pipe_in_a_declared_label_cannot_break_the_header_row():
    """One stray "|" in the HEADER breaks the whole table rather than its own row.
    prompt_capability's pattern already excludes it; this is the second lock."""
    out = _day_table_from_rows([{"day_number": 1}], ["Bad | Label"])
    width = len(_DAY_TABLE_HEADER) + 1
    assert len(out[0].strip("|").split("|")) == width
    assert out[1].count("---") == width


def test_the_renderer_and_the_fill_stage_use_one_missing_marker():
    """Two copies of this string would show a reviewer two different markers for the
    same condition."""
    out = _day_table_from_rows([{"day_number": 1}], ["Stage"])
    assert out[2].strip("|").split("|")[-1].strip() == EXTENSION_MISSING


# --------------------------------------------------------------------------- #
# Pass 4 must describe the table it was given, not the one it assumes
# --------------------------------------------------------------------------- #
def _cov_ok() -> dict:
    return {"total_days": 1, "days_in_output": 1, "missing_days": [], "declared_acs": ["A"],
            "covered_acs": ["A"], "orphan_acs": [], "failed_days": [], "thin_days": []}


def test_pass_four_names_the_blank_cells_instead_of_denying_them():
    """Found in review: it asserted "no cell is blank", which the renderer contradicts
    on every row — six columns default to "" by design."""
    table = _day_table_from_rows([{"day_number": 1, "topic": "x"}])
    line = _self_review_lines(_cov_ok(), 5, table)[5]
    assert "no cell is blank" not in line
    assert "AIM SME Comments (1)" in line and "Empty cells" in line


def test_pass_four_says_nothing_about_blanks_when_there_are_none():
    """A fully-populated row must not draw a warning about cells that are all filled."""
    row = {label.lower().replace(" ", "_"): "v" for label in _DAY_TABLE_HEADER}
    row.update({"day_number": 1, "topic": "t", "narrative": "n", "how_it_is_applied": "h",
                "concept_type_explanation": "e", "summative_exam_cluster": "c",
                "academian_questions": "q"})
    table = _day_table_from_rows([row], ["Stage"])
    # AIM SME Comments is reviewer-fill and always blank, so exactly one column remains
    line = _self_review_lines(_cov_ok(), 5, table)[5]
    assert "AIM SME Comments (1)" in line
    assert "Notes" not in line and "How It Is Applied" not in line


def test_pass_four_without_a_table_still_reports_the_flag_half():
    assert "recorded vocabulary" in _self_review_lines(_cov_ok(), 5)[5]
