"""AIM Hangar Activities workbook parser.

Parses multi-sheet hangar summary workbooks (e.g. "Hangar Activities Summary
Blocks 5-16.xlsx"):

* Block N (Day-Night|Weekend) sheets — day rows with Optional Hangar Activities
* Hangar Activities Summary — flat table used only to fill empty hangar cells
* ASA Reference — skipped

Each Block sheet becomes one ``hangar_sheet`` content unit. Day rows are not
written to ``dis_calendar_days`` (calendars own that table).
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Tuple

from services.aim_calendar import (
    _BLOCK_RE,
    _MAX_COLS,
    _banner_text,
    _block_number_from_text,
    _clean,
    _is_header,
    _is_punct,
    _is_subject_section_header,
    _schedule_from_name,
    _split_list,
    expand_acs,
)
from services.blocks import block_label

_SUMMARY_HEADERS = ("block tab", "hangar activity")
_HANGAR_COL_RE = re.compile(r"\bhangar\b", re.I)
_BLOCK_SHEET_RE = re.compile(
    r"^\s*block\s*0*(\d+)\s*\((?:day[\s\-_/]*night|weekend)\)\s*$",
    re.I,
)
_ASA_SKIP_RE = re.compile(r"^\s*asa\b", re.I)
_SUMMARY_SHEET_RE = re.compile(r"hangar\s+activities?\s+summary", re.I)


def _load_workbook(content: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content), data_only=True)


def _normalize_tab_key(name: str) -> str:
    """Normalize sheet / Block Tab labels for Summary → sheet matching."""
    s = re.sub(r"\s+", " ", (name or "").strip().lower())
    s = s.replace("day/night", "day-night").replace("day night", "day-night")
    return s


def _is_summary_header(cells: List[str]) -> bool:
    text = " ".join(c.strip().lower() for c in cells if c and str(c).strip())
    return all(h in text for h in _SUMMARY_HEADERS)


def _find_hangar_col(header_cells: List[str]) -> Optional[int]:
    """0-based index of the Optional Hangar Activities column."""
    for i, cell in enumerate(header_cells):
        if _HANGAR_COL_RE.search(cell or ""):
            return i
    # AIM layout: hangar is typically column K (index 10).
    if len(header_cells) > 10:
        return 10
    return None


def _sheet_has_aim_header(ws, max_scan: int = 8) -> bool:
    for row in ws.iter_rows(min_row=1, max_row=max_scan, max_col=_MAX_COLS, values_only=True):
        if _is_header([_clean(c) for c in row]):
            return True
    return False


def _sheet_has_summary_header(ws, max_scan: int = 5) -> bool:
    for row in ws.iter_rows(min_row=1, max_row=max_scan, max_col=6, values_only=True):
        if _is_summary_header([_clean(c) for c in row]):
            return True
    return False


def looks_like_hangar_activities_workbook(content: bytes) -> bool:
    """True for hangar summary / Block-sheet workbooks with hangar columns."""
    try:
        wb = _load_workbook(content)
    except Exception:
        return False
    for ws in wb.worksheets:
        title = ws.title or ""
        if _ASA_SKIP_RE.match(title):
            continue
        if _SUMMARY_SHEET_RE.search(title) and _sheet_has_summary_header(ws):
            return True
        if _BLOCK_SHEET_RE.match(title) and _sheet_has_aim_header(ws):
            return True
        if _sheet_has_summary_header(ws):
            return True
    return False


def _parse_summary_fill(ws) -> Dict[Tuple[str, int], str]:
    """Map (normalized Block Tab, class day) → hangar activity prose."""
    fill: Dict[Tuple[str, int], str] = {}
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row or 0, max_col=4, values_only=True))
    if not rows:
        return fill
    start = 0
    for i, row in enumerate(rows[:5]):
        cells = [_clean(c) for c in row]
        if _is_summary_header(cells):
            start = i + 1
            break
    for row in rows[start:]:
        cells = [_clean(c) for c in (list(row) + ["", "", "", ""])[:4]]
        tab, subject, day_cell, hangar = cells
        if not tab or not hangar:
            continue
        if not re.fullmatch(r"\d+", day_cell or ""):
            continue
        key = (_normalize_tab_key(tab), int(day_cell))
        # Prefer first non-empty; Summary should be authoritative when filling.
        if key not in fill:
            fill[key] = hangar
            if subject:
                fill[key] = hangar  # hangar text alone; subject used on Block sheets
    return fill


def _compose_hangar_day_text(day_number: int, subject: str, hangar_items: List[str]) -> str:
    hangar = "; ".join(hangar_items)
    subj = (subject or "").strip()
    if subj:
        return f"Day {day_number} — {subj}: {hangar}"
    return f"Day {day_number}: {hangar}"


def _parse_block_sheet(
    ws,
    sheet_index: int,
    *,
    summary_fill: Optional[Dict[Tuple[str, int], str]] = None,
    block_hint: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    """Parse one Block Day-Night/Weekend sheet into hangar-focused day records."""
    if not _sheet_has_aim_header(ws):
        return None

    max_c = min(ws.max_column or 1, max(_MAX_COLS, 12))
    grid = [[_clean(ws.cell(r, c).value) for c in range(1, max_c + 1)]
            for r in range(1, (ws.max_row or 0) + 1)]
    for mr in ws.merged_cells.ranges:
        if mr.min_col > max_c:
            continue
        for r in range(mr.min_row, mr.max_row + 1):
            if r == mr.min_row:
                continue
            for c in range(mr.min_col, min(mr.max_col, max_c) + 1):
                if not grid[r - 1][c - 1]:
                    grid[r - 1][c - 1] = ""

    hangar_col: Optional[int] = None
    for row in grid[:8]:
        cells = (row + [""] * max_c)[:max_c]
        if _is_header(cells):
            hangar_col = _find_hangar_col(cells)
            break
    if hangar_col is None:
        hangar_col = 10 if max_c > 10 else None

    banner = _banner_text(grid)
    block_num = _block_number_from_text(ws.title, banner, block_hint)
    schedule = _schedule_from_name(ws.title)
    tab_key = _normalize_tab_key(ws.title)

    raw_days: List[Dict[str, Any]] = []
    current: Optional[Dict[str, Any]] = None
    current_subject: Optional[str] = None
    for row in grid:
        cells = (row + [""] * max_c)[:max_c]
        if _is_header(cells):
            continue
        section = _is_subject_section_header(cells)
        if section:
            current_subject = section
            current = None
            continue
        day_cell = cells[0].strip()
        if re.fullmatch(r"\d+", day_cell):
            subject_cell = cells[1].strip() if len(cells) > 1 else ""
            if current_subject and re.fullmatch(r"\d+", subject_cell):
                cells = list(cells)
                cells[1] = f"{current_subject} - Day {subject_cell}"
            elif current_subject and not subject_cell:
                cells = list(cells)
                cells[1] = current_subject
            current = {"day_number": int(day_cell), "cells": [[c] if c else [] for c in cells]}
            raw_days.append(current)
        elif current:
            content_cells = [c for c in cells[1:] if c]
            if content_cells and not all(_is_punct(c) for c in content_cells):
                for i in range(1, max_c):
                    if cells[i] and not _is_punct(cells[i]):
                        current["cells"][i].append(cells[i])

    days: List[Dict[str, Any]] = []
    for rd in raw_days:
        joined = ["\n".join(x) for x in rd["cells"]]
        day_number = rd["day_number"]
        subject = joined[1] if len(joined) > 1 else ""
        hangar_items = (
            _split_list(joined[hangar_col])
            if hangar_col is not None and hangar_col < len(joined)
            else []
        )
        if not hangar_items and summary_fill:
            filled = summary_fill.get((tab_key, day_number))
            if filled:
                hangar_items = _split_list(filled)
        if not hangar_items:
            continue
        acs = expand_acs(joined[3] if len(joined) > 3 else "")
        project_acs = expand_acs(joined[6] if len(joined) > 6 else "")
        merged_acs: List[str] = []
        seen = set()
        for code in acs + project_acs:
            if code not in seen:
                seen.add(code)
                merged_acs.append(code)
        days.append({
            "day_number": day_number,
            "subject": subject,
            "hangar_activities": hangar_items,
            "acs_codes": merged_acs,
            "source_text": _compose_hangar_day_text(day_number, subject, hangar_items),
            "block": block_label(block_num) if block_num else "",
            "block_number": block_num,
            "block_id": f"B{block_num}" if block_num else "",
        })

    days.sort(key=lambda d: d["day_number"])
    if not days:
        return None
    return {
        "sheet_name": ws.title,
        "sheet_index": sheet_index,
        "block": block_label(block_num) if block_num else "",
        "block_number": block_num,
        "block_id": f"B{block_num}" if block_num else "",
        "schedule": schedule,
        "banner": banner[:500] if banner else "",
        "total_days_with_hangar": len(days),
        "total_days_detected": len(days),
        "days": days,
    }


def build_hangar_structure(
    content: bytes,
    filename: str,
    block_hint: Optional[str] = None,
) -> Dict[str, Any]:
    """Parse a hangar activities workbook into ``hangar_structure``."""
    wb = _load_workbook(content)
    file_block = _block_number_from_text(filename)
    hint_block = _block_number_from_text(block_hint) if block_hint else None
    effective_hint = block_label(hint_block or file_block) if (hint_block or file_block) else None

    summary_fill: Dict[Tuple[str, int], str] = {}
    for ws in wb.worksheets:
        if _SUMMARY_SHEET_RE.search(ws.title or "") or _sheet_has_summary_header(ws):
            if not _ASA_SKIP_RE.match(ws.title or ""):
                summary_fill.update(_parse_summary_fill(ws))

    sheets: List[Dict[str, Any]] = []
    skipped: List[Dict[str, str]] = []
    for idx, ws in enumerate(wb.worksheets):
        title = ws.title or f"Sheet{idx}"
        if _ASA_SKIP_RE.match(title):
            skipped.append({"sheet": title, "reason": "asa_reference_skipped"})
            continue
        if _SUMMARY_SHEET_RE.search(title) or (
            _sheet_has_summary_header(ws) and not _BLOCK_SHEET_RE.match(title)
        ):
            skipped.append({"sheet": title, "reason": "summary_used_as_fill_only"})
            continue
        if not _BLOCK_SHEET_RE.match(title) and not _sheet_has_aim_header(ws):
            skipped.append({"sheet": title, "reason": "not_a_block_hangar_sheet"})
            continue
        try:
            parsed = _parse_block_sheet(
                ws, idx, summary_fill=summary_fill, block_hint=effective_hint,
            )
        except Exception as exc:  # noqa: BLE001
            skipped.append({"sheet": title, "reason": f"{type(exc).__name__}: {exc}"})
            continue
        if parsed is None:
            skipped.append({"sheet": title, "reason": "no_hangar_day_rows"})
            continue
        sheets.append(parsed)

    blocks_covered: List[str] = []
    for s in sheets:
        b = s.get("block")
        if b and b not in blocks_covered:
            blocks_covered.append(b)

    return {
        "structure_type": "hangar_activities_workbook",
        "processing_profile": "aim_hangar_activities",
        "profile": "aim_hangar_activities",
        "blocks_covered": blocks_covered,
        "total_sheets_detected": len(sheets),
        "sheets": sheets,
        "skipped_sheets": skipped,
        "confidence": "high" if sheets else "low",
        "warnings": [] if sheets else ["No hangar activity day rows detected."],
    }


def iter_hangar_sheets(hangar_structure: Dict[str, Any]) -> List[Dict[str, Any]]:
    if not hangar_structure:
        return []
    return list(hangar_structure.get("sheets") or [])
