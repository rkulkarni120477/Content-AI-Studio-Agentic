"""AKTR knowledge-test report parser — one record per block, per ACS code.

An AKTR (Airman Knowledge Test Report) rollup is one workbook covering the whole
program: a sheet per block, each captioned "ALL CAMPUS — Block N | Top 10 Most
Missed ACS Codes" over four columns — Rank, % Missed, ACS Code, ACS Code
Description. It is the only source of real knowledge-test performance data, which
a block-wide Blueprint needs to fill "Targets for Quick Check" and the per-code
high-miss field with figures instead of a placeholder.

Two properties make it unlike every other AIM source document:

* It is **not one block's document.** Word-chunking it and stamping the whole file
  with a single ``block`` (``_extract_block_day`` reads the first 1,500 characters,
  so it sees only the FIRST sheet) files all sixteen blocks' data under Block 1,
  where the block-scoped digest loader — ``metadata_json->>'block' = %s`` in
  ``services/digests/profiles/aim.py`` — can never find it for any other block.
* Its meaning is per row, keyed by ACS code. A chunk boundary landing mid-sheet
  splits a block's table across two units.

So this module reads the workbook directly and emits one record per block sheet,
which ``content_unit_creation_agent`` turns into one content unit per block, each
carrying that block's own ``block``/``block_number`` metadata and its ``acs_codes``.

Entry points:
    looks_like_aktr_report(content) -> bool   # self-gating format check
    build_knowledge_test_structure(content, filename) -> dict
    parse_missed_codes(text) -> dict          # code -> {pct_missed, rank, description}
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional

#: An ACS code as AIM writes it: AM.II.K.K1, AM.III.F.K10.
ACS_CODE_RE = re.compile(r"\bAM\.[IVX]+\.[A-Z]\.[A-Z]\d+\b")

#: Header cells that identify an AKTR sheet. Matched case-insensitively against the
#: header row's own names, so a reordered or extra column changes nothing.
_REQUIRED_HEADERS = {"acs code", "% missed"}

#: Caption/sheet-name form that carries the block number.
_BLOCK_RE = re.compile(r"\bblock\s*0*(\d+)\b", re.I)

#: Sheets a rollup can carry that are not per-block tables.
_SKIP_SHEET_NAMES = {"summary", "notes", "readme", "cover", "index", "all campus"}


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header_map(row: List[Any]) -> Dict[str, int]:
    """Column index per normalized header name."""
    return {_clean(c).lower(): i for i, c in enumerate(row) if _clean(c)}


def _find_header(rows: List[tuple], scan_limit: int = 10) -> Optional[int]:
    """Index of the AKTR header row, or None when this sheet has none.

    Mirrors ``extractors._xlsx_header_row``: the caption above the header is a
    merged single-cell banner, so the header is not row 1.
    """
    for idx, row in enumerate(rows[:scan_limit]):
        names = set(_header_map(list(row)))
        if _REQUIRED_HEADERS <= names:
            return idx
    return None


def _pct(raw: str) -> Optional[float]:
    """A miss rate as a fraction of 1.

    The workbook stores 0.4816753926701571 for 48.2%; a hand-edited sheet may hold
    "48.2%" or "48.2" instead. Anything above 1 is read as already-percent so a
    hand-edited sheet does not report 4,816%.
    """
    text = raw.replace("%", "").strip()
    if not text:
        return None
    try:
        value = float(text)
    except ValueError:
        return None
    if value > 1:
        value = value / 100.0
    if not 0 <= value <= 1:
        return None
    return value


def _block_number(sheet_name: str, caption: str) -> Optional[int]:
    """Block number from the sheet name, falling back to its caption."""
    for text in (sheet_name, caption):
        m = _BLOCK_RE.search(text or "")
        if m:
            return int(m.group(1))
    return None


def _load_rows(content: bytes) -> List[tuple]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out: List[tuple] = []
    for name in wb.sheetnames:
        out.append((name, list(wb[name].iter_rows(values_only=True))))
    return out


def looks_like_aktr_report(content: bytes) -> bool:
    """Whether this workbook is an AKTR missed-code rollup.

    Self-gating on structure, not on filename, so a renamed file still parses and a
    same-named file of some other shape falls through to the generic path.
    """
    try:
        sheets = _load_rows(content)
    except Exception:
        return False
    for _name, rows in sheets:
        if rows and _find_header(rows) is not None:
            return True
    return False


def build_knowledge_test_structure(content: bytes, filename: str = "") -> Dict[str, Any]:
    """One record per block sheet, in the shape ``content_unit_creation_agent`` reads.

    Sheets without an AKTR header row, and sheets whose block number cannot be
    determined, are reported in ``skipped_sheets`` rather than dropped quietly — a
    block missing from the Blueprint's performance data must be traceable to the
    sheet it should have come from.
    """
    blocks: List[Dict[str, Any]] = []
    skipped: List[Dict[str, str]] = []
    try:
        sheets = _load_rows(content)
    except Exception as exc:  # noqa: BLE001 — caller falls back to generic chunking
        return {"structure_type": "knowledge_test_report", "blocks": [],
                "skipped_sheets": [{"sheet": "*", "reason": f"unreadable workbook: {exc}"}],
                "source_file_name": filename}

    for name, rows in sheets:
        if not rows:
            skipped.append({"sheet": name, "reason": "empty sheet"})
            continue
        header_index = _find_header(rows)
        if header_index is None:
            if name.strip().lower() not in _SKIP_SHEET_NAMES:
                skipped.append({"sheet": name, "reason": "no AKTR header row (Rank / % Missed / ACS Code)"})
            continue
        caption = " ".join(_clean(c) for r in rows[:header_index] for c in r if _clean(c))
        block_number = _block_number(name, caption)
        if not block_number:
            skipped.append({"sheet": name, "reason": "no block number in sheet name or caption"})
            continue

        cols = _header_map(list(rows[header_index]))
        code_col = cols.get("acs code")
        pct_col = cols.get("% missed")
        rank_col = cols.get("rank")
        desc_col = next((i for n, i in cols.items() if "description" in n), None)

        codes: List[Dict[str, Any]] = []
        for row in rows[header_index + 1:]:
            cells = [_clean(c) for c in row]
            if not any(cells):
                continue
            code = cells[code_col] if code_col is not None and code_col < len(cells) else ""
            if not ACS_CODE_RE.fullmatch(code):
                continue
            entry: Dict[str, Any] = {"acs_code": code}
            if pct_col is not None and pct_col < len(cells):
                pct = _pct(cells[pct_col])
                if pct is not None:
                    entry["pct_missed"] = pct
            if rank_col is not None and rank_col < len(cells) and cells[rank_col].isdigit():
                entry["rank"] = int(cells[rank_col])
            if desc_col is not None and desc_col < len(cells):
                entry["description"] = cells[desc_col]
            codes.append(entry)

        if not codes:
            skipped.append({"sheet": name, "reason": "header row present but no ACS-code rows"})
            continue
        blocks.append({
            "block": f"Block {block_number}",
            "block_number": block_number,
            "sheet": name,
            "caption": caption,
            "codes": codes,
            "source_text": _compose_text(block_number, caption, codes),
        })

    return {"structure_type": "knowledge_test_report", "blocks": blocks,
            "skipped_sheets": skipped, "source_file_name": filename}


def _compose_text(block_number: int, caption: str, codes: List[Dict[str, Any]]) -> str:
    """The unit's embedded/retrieved text: a readable table, percentages spelled out.

    The raw workbook value is a float fraction (0.4816753926701571). Left as-is it
    reads as neither a percentage nor a rank to a model, so each row states the
    figure the way the Blueprint has to print it.
    """
    lines = [caption or f"Block {block_number} — Most Missed ACS Codes (AKTR knowledge-test data)"]
    for entry in codes:
        rank = entry.get("rank")
        pct = entry.get("pct_missed")
        parts = [f"#{rank}" if rank else "", entry["acs_code"]]
        if pct is not None:
            parts.append(f"{pct * 100:.1f}% missed")
        if entry.get("description"):
            parts.append(entry["description"])
        lines.append(" — ".join(p for p in parts if p))
    return "\n".join(lines)


def parse_missed_codes(text: str) -> Dict[str, Dict[str, Any]]:
    """Read back a unit's ``source_text`` as ``code -> {pct_missed, rank, description}``.

    The worksheet builder reads the structure store, which holds the composed text
    rather than the workbook, so the same format is parsed back here. Kept beside
    ``_compose_text`` so the two cannot drift.
    """
    out: Dict[str, Dict[str, Any]] = {}
    for line in (text or "").splitlines():
        m = ACS_CODE_RE.search(line)
        if not m:
            continue
        code = m.group(0)
        entry: Dict[str, Any] = {}
        rank = re.search(r"#(\d+)\s*—", line)
        if rank:
            entry["rank"] = int(rank.group(1))
        tail = line.split(code, 1)[1]
        pct = re.search(r"([\d.]+)%\s*missed", tail, re.I)
        if pct:
            try:
                entry["pct_missed"] = float(pct.group(1)) / 100.0
            except ValueError:
                pass
        # Everything after the last field this format defines is the description.
        # Taken as one remainder rather than by splitting on the em-dash separator,
        # so a description that itself contains an em dash survives intact.
        description = (tail[pct.end():] if pct else tail).lstrip(" —-").strip()
        if description:
            entry["description"] = description
        # First mention wins: a code repeated across sheets belongs to the block whose
        # unit is being read, and the caller reads one block's units.
        out.setdefault(code, entry)
    return out
