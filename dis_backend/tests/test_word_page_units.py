"""Word uploads become page units (ebook-equivalent), not word-chunks."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_BREAK

from services.agents.content_unit_creation_agent import ContentUnitCreationAgent
from services.ebook_page_retag import merge_retag_units_into_list
from services.ebook_page_tagger import TAGGING_OK, TAGGING_PENDING
from services.source_library import build_clean_content_document, chunking_strategy


class _FakeCtx:
    def __init__(self):
        from config.settings import get_tenant_config
        self.cfg = get_tenant_config("aim")
        self.cfg.pipeline.llm_provider = "mock"
        self.cfg.pipeline.bedrock_enabled = False
        self.guard = None

    def step_done(self, state, name):
        return state


def _agent():
    agent = ContentUnitCreationAgent.__new__(ContentUnitCreationAgent)
    agent.ctx = _FakeCtx()
    return agent


def _docx_bytes(*paragraphs: str, page_break_after: int | None = None) -> bytes:
    doc = Document()
    for i, text in enumerate(paragraphs):
        doc.add_paragraph(text)
        if page_break_after is not None and i == page_break_after:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_quiz_produces_page_units():
    raw = _docx_bytes("Q1 landing gear", "Q2 torque link", page_break_after=0)
    agent = _agent()
    state = {
        "job_id": "wq1",
        "doc_type": "quiz",
        "filename": "B2Q1.docx",
        "file_type": "docx",
        "doc_metadata": {
            "document_type": "quiz",
            "content_type": "quiz",
            "title": "Block 2 Quiz 1",
            "purpose": "course_generation",
        },
        "raw_bytes": raw,
        "raw_text": "",
        "slide_texts": [],
    }
    units = agent.run(state)["content_units"]
    assert len(units) == 2
    assert units[0]["content_unit_id"] == "wq1:page_1"
    assert units[0]["unit_type"] == "page"
    assert units[0]["metadata"]["chunking_strategy"] == "page"
    assert units[0]["metadata"]["page_number"] == "1"
    assert units[0]["metadata"]["tagging_status"] == TAGGING_PENDING
    assert "Q1" in units[0]["text"]
    assert "Q2" in units[1]["text"]


def test_word_no_break_is_single_page_unit():
    raw = _docx_bytes("Entire exam on one page with no breaks.")
    agent = _agent()
    state = {
        "job_id": "wq2",
        "doc_type": "final_exam",
        "filename": "Block_05_Final.docx",
        "file_type": "docx",
        "doc_metadata": {
            "document_type": "final_exam",
            "content_type": "final_exam",
            "title": "Final Exam",
            "purpose": "course_generation",
        },
        "raw_bytes": raw,
        "raw_text": "",
        "slide_texts": [],
    }
    units = agent.run(state)["content_units"]
    assert len(units) == 1
    assert units[0]["content_unit_id"] == "wq2:page_1"
    assert units[0]["unit_type"] == "page"


def test_word_instructor_guide_is_page_not_full_document():
    """Style/CDD/blueprint Word files used to collapse to one welded unit."""
    raw = _docx_bytes("Guide page 1", "Guide page 2", page_break_after=0)
    agent = _agent()
    state = {
        "job_id": "wg1",
        "doc_type": "instructor_guide",
        "filename": "Instructor_Guide.docx",
        "file_type": "docx",
        "doc_metadata": {
            "document_type": "instructor_guide",
            "content_type": "instructor_guide",
            "title": "IG",
            "purpose": "cdd",
            "use_for_cdd": True,
        },
        "raw_bytes": raw,
        "raw_text": "",
        "slide_texts": [],
    }
    units = agent.run(state)["content_units"]
    assert len(units) == 2
    assert all(u["unit_type"] == "page" for u in units)


def test_hangar_word_uses_hangar_page_units():
    raw = _docx_bytes("Hangar day 1", "Hangar day 2", page_break_after=0)
    agent = _agent()
    state = {
        "job_id": "wh1",
        "doc_type": "hangar_activity",
        "filename": "Hangar_Block5.docx",
        "file_type": "docx",
        "doc_metadata": {
            "document_type": "hangar_activity",
            "content_type": "hangar_activity",
            "title": "Hangar",
        },
        "raw_bytes": raw,
        "raw_text": "",
        "slide_texts": [],
        "hangar_structure": {},  # not an xlsx workbook
    }
    units = agent.run(state)["content_units"]
    assert len(units) == 2
    assert units[0]["unit_type"] == "hangar_page"
    assert units[0]["metadata"]["page_number"] == "1"


def test_lesson_pdf_still_word_chunks():
    agent = _agent()
    words = " ".join(f"w{i}" for i in range(600))
    state = {
        "job_id": "j2",
        "doc_type": "lesson_pdf",
        "filename": "lesson.pdf",
        "file_type": "pdf",
        "doc_metadata": {"document_type": "lesson_pdf", "content_type": "lesson_pdf"},
        "raw_text": words,
        "slide_texts": [],
    }
    units = agent.run(state)["content_units"]
    assert units[0]["content_unit_id"].startswith("j2:unit_")


def test_chunking_strategy_page_from_unit_type():
    units = [{"unit_type": "page", "text": "x", "metadata": {"page_number": "1"}}]
    assert chunking_strategy("course_generation", "quiz", units=units) == "page"
    assert chunking_strategy("course_generation", "final_exam", units=units) == "page"
    assert chunking_strategy("cdd", "instructor_guide", units=units) == "page"


def test_content_json_keeps_word_page_metadata():
    doc = build_clean_content_document({
        "job_id": "j",
        "metadata": {"document_type": "quiz", "title": "Q1", "purpose": "course_generation"},
        "source_file": {"name": "B2Q1.docx", "type": "docx"},
        "content_units": [{
            "content_unit_id": "j:page_1",
            "unit_type": "page",
            "unit_number": 1,
            "title": "Q1 — p. 1",
            "text": "Question text",
            "topics": ["landing gear"],
            "metadata": {
                "pdf_page": 1,
                "page_number": "1",
                "chunking_strategy": "page",
                "tagging_status": "ok",
                "acs_codes": ["AM.I.D.K1"],
                "summary": "Landing gear intro",
                "raw_s3_key": "should-strip",
            },
        }],
    })
    assert doc["chunking_strategy"] == "page"
    unit = doc["content_units"][0]
    assert unit["metadata"]["page_number"] == "1"
    assert unit["metadata"]["acs_codes"] == ["AM.I.D.K1"]
    assert "raw_s3_key" not in unit["metadata"]


def test_retag_merge_preserves_sibling_word_pages():
    units = [
        {
            "content_unit_id": "j:page_1",
            "title": "p1",
            "topics": ["oleo"],
            "metadata": {"tagging_status": TAGGING_OK, "summary": "A", "page_number": "1"},
        },
        {
            "content_unit_id": "j:page_2",
            "title": "p2",
            "topics": [],
            "metadata": {"tagging_status": TAGGING_PENDING, "page_number": "2"},
        },
    ]
    updated = {
        "j:page_2": {
            "content_unit_id": "j:page_2",
            "title": "p2 tagged",
            "topics": ["brakes"],
            "metadata": {
                "tagging_status": TAGGING_OK,
                "summary": "Brakes.",
                "page_number": "2",
                "topics": ["brakes"],
            },
        },
    }
    merge_retag_units_into_list(units, updated, replace_metadata=False)
    assert units[0]["metadata"]["tagging_status"] == TAGGING_OK
    assert units[0]["metadata"]["summary"] == "A"
    assert units[1]["metadata"]["tagging_status"] == TAGGING_OK
    assert units[1]["topics"] == ["brakes"]
