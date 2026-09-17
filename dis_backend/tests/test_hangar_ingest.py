"""Hangar activities workbook ingest — Block sheets + Summary fill."""
from __future__ import annotations

import io
from typing import Any, Dict, List

import openpyxl
import pytest

from services.aim_hangar import (
    build_hangar_structure,
    iter_hangar_sheets,
    looks_like_hangar_activities_workbook,
)
from services.indexing import _build_bulk_actions
from services.source_library import build_clean_content_document, chunking_strategy


def _write_block_sheet(
    ws,
    *,
    banner: str,
    subject: str,
    days: List[Dict[str, Any]],
    hangar_col: bool = True,
):
    ws["B1"] = banner
    headers = [
        "Days", "Subject/Day", "Topics Covered", "ACS Codes for Topics",
        "Corresponding Handbook Pages", "Projects", "ACS Codes for Projects",
        "Quiz", "Supplemental Resources", "Test Prep Activities",
        "Optional Hangar Activities", "Additional Reading",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(2, col, h)
    ws.cell(3, 2, subject)
    row = 4
    for day in days:
        ws.cell(row, 1, day["day"])
        ws.cell(row, 2, day.get("subj_day", day["day"]))
        ws.cell(row, 3, day.get("topics", "Topic A"))
        ws.cell(row, 4, day.get("acs", "AM.I.B.K1"))
        if hangar_col:
            ws.cell(row, 11, day.get("hangar", ""))
        row += 1


def _workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()

    summary = wb.active
    summary.title = "Hangar Activities Summary"
    summary["A1"] = "Block Tab"
    summary["B1"] = "Subject"
    summary["C1"] = "Class Day"
    summary["D1"] = "Hangar Activity"
    summary["A2"] = "Block 5 (Day-Night)"
    summary["B2"] = "Metallic Structures"
    summary["C2"] = 1
    summary["D2"] = "Summary fill: show structure examples"

    asa = wb.create_sheet("ASA Reference")
    asa["A1"] = "ASA Textbook Reference Guide"

    ws5 = wb.create_sheet("Block 5 (Day-Night)")
    _write_block_sheet(
        ws5,
        banner="Block 5 (Day/Night): Metallic Structures",
        subject="Metallic Structures",
        days=[
            # Empty hangar — should fill from Summary
            {"day": 1, "subj_day": 1, "topics": "Alloys", "acs": "AM.I.B.K1", "hangar": ""},
            {"day": 2, "subj_day": 2, "topics": "Drilling", "acs": "AM.I.B.K2",
             "hangar": "Instructor will demonstrate drilling techniques"},
        ],
    )

    ws1 = wb.create_sheet("Block 1 (Weekend)")
    _write_block_sheet(
        ws1,
        banner="Block 1 (Weekend): Fundamentals of Math",
        subject="Fundamentals of Math",
        days=[
            {"day": 1, "subj_day": 1, "topics": "Measuring", "acs": "AM.I.H.K1",
             "hangar": "Measuring the Area of Shop Floor"},
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_looks_like_hangar_workbook():
    assert looks_like_hangar_activities_workbook(_workbook_bytes())


def test_build_hangar_structure_skips_asa_and_summary_as_section():
    result = build_hangar_structure(
        _workbook_bytes(), "Hangar Activities Summary Blocks 5-16.xlsx",
    )
    assert result["structure_type"] == "hangar_activities_workbook"
    sheets = iter_hangar_sheets(result)
    names = [s["sheet_name"] for s in sheets]
    assert "ASA Reference" not in names
    assert "Hangar Activities Summary" not in names
    assert "Block 5 (Day-Night)" in names
    assert "Block 1 (Weekend)" in names
    assert "Block 5" in result["blocks_covered"]
    assert "Block 1" in result["blocks_covered"]

    reasons = {s["sheet"]: s["reason"] for s in result["skipped_sheets"]}
    assert "asa_reference_skipped" in reasons.get("ASA Reference", "")
    assert "summary_used_as_fill_only" in reasons.get("Hangar Activities Summary", "")


def test_summary_fills_empty_hangar_cell():
    result = build_hangar_structure(_workbook_bytes(), "Hangar Activities Summary.xlsx")
    sheets = {s["sheet_name"]: s for s in iter_hangar_sheets(result)}
    b5 = sheets["Block 5 (Day-Night)"]
    day1 = next(d for d in b5["days"] if d["day_number"] == 1)
    assert "Summary fill" in day1["source_text"]
    day2 = next(d for d in b5["days"] if d["day_number"] == 2)
    assert "drilling" in day2["source_text"].lower()


def test_weekend_schedule_and_block_stamp():
    result = build_hangar_structure(_workbook_bytes(), "hangar.xlsx")
    sheets = {s["sheet_name"]: s for s in iter_hangar_sheets(result)}
    we = sheets["Block 1 (Weekend)"]
    assert we["schedule"] == "weekend"
    assert we["block"] == "Block 1"
    assert we["block_number"] == 1
    assert we["days"][0]["hangar_activities"]


def test_source_library_per_unit_for_hangar_sheets():
    units = [{
        "content_unit_id": "j:hangar_sheet_0",
        "unit_type": "hangar_sheet",
        "unit_number": 1,
        "title": "Block 5 (Day-Night)",
        "text": "Day 1 — Metallic Structures: demo",
        "metadata": {
            "block": "Block 5",
            "sheet_name": "Block 5 (Day-Night)",
            "sheet_index": 0,
            "schedule": "day_night",
            "document_type": "hangar_activity",
            "tagging_status": "pending",
        },
    }]
    assert chunking_strategy("course_generation", "hangar_activity", units=units) == "per_unit"
    doc = build_clean_content_document({
        "job_id": "j",
        "metadata": {"document_type": "hangar_activity", "title": "Hangar Summary",
                     "blocks_covered": ["Block 5"]},
        "source_file": {"name": "Hangar Activities Summary.xlsx", "type": "xlsx"},
        "content_units": units,
    })
    assert doc["chunking_strategy"] == "per_unit"
    assert doc["content_units"][0]["unit_type"] == "hangar_sheet"
    assert doc["content_units"][0]["metadata"]["block"] == "Block 5"
    assert doc["tagging_pending_count"] == 1


def test_opensearch_bulk_uses_hangar_activity_type():
    state = {
        "job_id": "j1",
        "tenant_id": "t",
        "client_id": "aim",
        "filename": "Hangar Activities Summary.xlsx",
        "file_type": "xlsx",
        "doc_type": "project_activity",
        "doc_metadata": {"content_type": "hangar_activity", "document_type": "hangar_activity"},
    }
    units = [{
        "content_unit_id": "j1:hangar_sheet_0",
        "unit_type": "hangar_sheet",
        "unit_number": 1,
        "title": "Block 5 (Day-Night)",
        "text": "Day 1 hangar text",
        "keywords": [],
        "topics": [],
        "metadata": {
            "block": "Block 5",
            "document_type": "hangar_activity",
            "content_type": "hangar_activity",
        },
        "embedding": [0.1, 0.2],
    }]
    actions = _build_bulk_actions("dis-content-dev-aim", state, units)
    assert actions[0]["_source"]["document_type"] == "hangar_activity"
    assert actions[0]["_source"]["block"] == "Block 5"


class _FakeCtx:
    def __init__(self):
        from config.settings import get_tenant_config
        self.cfg = get_tenant_config("aim")
        self.cfg.pipeline.llm_provider = "mock"
        self.cfg.pipeline.bedrock_enabled = False
        self.guard = None

    def step_done(self, state, name):
        return state


def test_content_unit_agent_hangar_sheets():
    from services.agents.content_unit_creation_agent import ContentUnitCreationAgent

    agent = ContentUnitCreationAgent.__new__(ContentUnitCreationAgent)
    agent.ctx = _FakeCtx()
    hangar = build_hangar_structure(_workbook_bytes(), "Hangar Activities Summary.xlsx")
    state = {
        "job_id": "hj",
        "doc_type": "project_activity",
        "filename": "Hangar Activities Summary.xlsx",
        "file_type": "xlsx",
        "doc_metadata": {
            "content_type": "hangar_activity",
            "document_type": "hangar_activity",
            "blocks_covered": hangar["blocks_covered"],
        },
        "hangar_structure": hangar,
        "raw_text": "",
    }
    units = agent.run(state)["content_units"]
    assert units
    assert all(u["unit_type"] == "hangar_sheet" for u in units)
    assert units[0]["title"] in {"Block 5 (Day-Night)", "Block 1 (Weekend)"}
    assert units[0]["metadata"]["document_type"] == "hangar_activity"
    assert units[0]["metadata"].get("tagging_status") == "pending"


def test_content_unit_agent_hangar_pdf_pages():
    from services.agents.content_unit_creation_agent import ContentUnitCreationAgent

    agent = ContentUnitCreationAgent.__new__(ContentUnitCreationAgent)
    agent.ctx = _FakeCtx()
    state = {
        "job_id": "hp",
        "doc_type": "project_activity",
        "filename": "B5D3 Hangar Lab.pdf",
        "file_type": "pdf",
        "doc_metadata": {
            "content_type": "hangar_activity",
            "document_type": "hangar_activity",
            "title": "B5D3 Hangar Lab",
            "block": "Block 5",
            "block_number": 5,
            "day_number": 3,
        },
        "page_texts": [
            {"pdf_page": 1, "text": "Shop safety briefing"},
            {"pdf_page": 2, "text": "Torque wrench procedure"},
        ],
        "page_count": 2,
        "raw_text": "",
        "raw_bytes": b"",
    }
    units = agent.run(state)["content_units"]
    assert len(units) == 2
    assert all(u["unit_type"] == "hangar_page" for u in units)
    assert units[0]["metadata"]["block"] == "Block 5"
    assert units[0]["metadata"]["day_number"] == 3
    assert units[0]["metadata"]["page_number"] == "1"
    assert units[0]["metadata"]["document_type"] == "hangar_activity"
