"""Word page-break splitting for ebook-equivalent page units.

Pins extract_docx_pages: hard page breaks, next-page section breaks,
lastRenderedPageBreak, no-break → 1 unit, and textbox text staying on
the owning page.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.enum.text import WD_BREAK

from services.pipeline.extractors import extract_docx, extract_docx_pages


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_zip(**parts: str | bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>",
        )
        for name, payload in parts.items():
            zf.writestr(name.replace("__", "/"), payload)
    return buf.getvalue()


def _python_docx_with_page_break(page1: str, page2: str) -> bytes:
    doc = Document()
    doc.add_paragraph(page1)
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)
    doc.add_paragraph(page2)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_hard_page_break_splits_into_two_pages():
    content = _python_docx_with_page_break(
        "Question 1: Identify the landing gear.",
        "Question 2: Torque link purpose.",
    )
    pages = extract_docx_pages(content)
    assert len(pages) == 2
    assert pages[0]["pdf_page"] == 1
    assert pages[1]["pdf_page"] == 2
    assert "Question 1" in pages[0]["text"]
    assert "Question 2" in pages[1]["text"]
    assert "Question 2" not in pages[0]["text"]


def test_no_page_break_yields_single_page():
    doc = Document()
    doc.add_paragraph("Single page quiz with no breaks.")
    buf = io.BytesIO()
    doc.save(buf)
    pages = extract_docx_pages(buf.getvalue())
    assert len(pages) == 1
    assert pages[0]["pdf_page"] == 1
    assert "Single page quiz" in pages[0]["text"]


def test_last_rendered_page_break_splits():
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body>
    <w:p><w:r><w:t>Page one content</w:t></w:r></w:p>
    <w:p>
      <w:r>
        <w:lastRenderedPageBreak/>
        <w:t>Page two content</w:t>
      </w:r>
    </w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""
    content = _docx_zip(**{"word/document.xml": document_xml})
    pages = extract_docx_pages(content)
    assert len(pages) == 2
    assert "Page one" in pages[0]["text"]
    assert "Page two" in pages[1]["text"]


def test_next_page_section_break_splits_after_paragraph():
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body>
    <w:p>
      <w:pPr>
        <w:sectPr>
          <w:type w:val="nextPage"/>
        </w:sectPr>
      </w:pPr>
      <w:r><w:t>End of section one</w:t></w:r>
    </w:p>
    <w:p><w:r><w:t>Start of section two</w:t></w:r></w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""
    content = _docx_zip(**{"word/document.xml": document_xml})
    pages = extract_docx_pages(content)
    assert len(pages) == 2
    assert "End of section one" in pages[0]["text"]
    assert "Start of section two" in pages[1]["text"]


def test_textbox_stays_on_owning_page():
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body>
    <w:p>
      <w:r>
        <w:drawing>
          <w:txbxContent>
            <w:p><w:r><w:t>Question 1 in textbox on page 1</w:t></w:r></w:p>
          </w:txbxContent>
        </w:drawing>
      </w:r>
    </w:p>
    <w:p>
      <w:r>
        <w:br w:type="page"/>
      </w:r>
    </w:p>
    <w:p><w:r><w:t>Question 2 on page 2</w:t></w:r></w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""
    content = _docx_zip(**{"word/document.xml": document_xml})
    pages = extract_docx_pages(content)
    assert len(pages) == 2
    assert "Question 1 in textbox on page 1" in pages[0]["text"]
    assert "Question 2 on page 2" in pages[1]["text"]
    assert "Question 2" not in pages[0]["text"]


def test_extract_docx_populates_page_texts():
    content = _python_docx_with_page_break("Alpha page", "Beta page")
    result = extract_docx(content)
    assert result.page_count == 2
    assert len(result.page_texts) == 2
    assert "Alpha" in result.page_texts[0]["text"]
    assert "Beta" in result.page_texts[1]["text"]


def test_legacy_ole2_doc_is_single_page():
    # Minimal OLE2 magic — extract falls through to scrape (empty) or one page.
    content = b"\xd0\xcf\x11\xe0" + b"\x00" * 64 + b"Question text here enough"
    pages = extract_docx_pages(content)
    # Scrape may or may not recover text; never more than one page for OLE2.
    assert len(pages) <= 1
