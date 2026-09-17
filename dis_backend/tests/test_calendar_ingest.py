"""Multi-sheet AIM calendar ingest + NEW FORMAT parse fixes.

Covers the ALL-blocks workbook shape: every sheet parsed, block from sheet
name (not a Windows "(1)" filename), sheet-scoped Postgres ids, and Source
Library per_unit sections titled with sheet names.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List

import openpyxl
import pytest

from services.aim_calendar import (
    _is_header,
    build_calendar_structure,
    expand_acs,
    iter_calendar_sheets,
    looks_like_aim_teacher_calendar,
    parse_handbook,
)
from services.indexing import upsert_calendar
from services.source_library import (
    PER_UNIT_DOC_TYPES,
    build_clean_content_document,
    chunking_strategy,
)


#: The 2026-08 workbooks (Block 9-15 "Calendar Day and Night.xlsx").
NEW_HEADER = ["Block Days", "Subject Days", "Topics Covered", "ACS Codes for Topics",
              "Corresponding Handbook", "Projects", "ACS Codes for Projects", "Quiz"]
#: The original layout the parser was written against.
OLD_HEADER = ["Days", "Subject/Day", "Topics Covered", "ACS Codes", "Handbook"]


@pytest.mark.parametrize("header", [NEW_HEADER, OLD_HEADER])
def test_both_calendar_wordings_are_recognised(header):
    assert _is_header(header)


@pytest.mark.parametrize("row", [
    ["1", "1", "Turbine operating principles", "AM.III.B.K1", "FAA-H-8083-32B"],
    ["", "", ""],
    ["Week", "Module", "Readings", "Assignments"],
    ["Day", "Notes"],
])
def test_non_headers_are_not_mistaken_for_one(row):
    assert not _is_header(row)


# ---------------------------------------------------------------------------
# Workbook fixtures
# ---------------------------------------------------------------------------

def _write_sheet(ws, *, banner: str, subject: str, days: List[Dict[str, Any]]):
    """NEW FORMAT layout: banner, header, subject section row, then day rows."""
    ws["B1"] = banner
    headers = [
        "Block Days", "Subject Days", "Topics Covered", "ACS Codes for Topics",
        "Corresponding Handbook Pages", "Projects", "ACS Codes for Projects",
        "Quiz", "Supplemental Resources", "Additional Reading",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(2, col, h)
    ws.cell(3, 2, subject)
    row = 4
    for day in days:
        ws.cell(row, 1, day["day"])
        ws.cell(row, 2, day.get("subj_day", day["day"]))
        ws.cell(row, 3, day.get("topics", "Topic A"))
        ws.cell(row, 4, day.get("acs", "AM.I.B.K1"))
        ws.cell(row, 5, day.get("handbook", "FAA-H-8083-30B • Ch. 4 pp. 4-1 to 4-9"))
        ws.cell(row, 6, day.get("projects", ""))
        ws.cell(row, 7, day.get("project_acs", ""))
        ws.cell(row, 8, day.get("quiz", ""))
        row += 1


def _two_sheet_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Block 1 (Day-Night)"
    _write_sheet(
        ws1,
        banner="Block 1 (Day/Night): Fundamentals of Math",
        subject="Fundamentals of Math",
        days=[
            {"day": 1, "subj_day": 1, "topics": "Fractions; Ratios",
             "acs": "AM.I.H.K3, K4, R3",
             "handbook": "FAA-H-8083-30B • Ch. 3 pp. 3-1 to 3-11",
             "projects": "Project 1-1"},
            {"day": 2, "subj_day": 2, "topics": "Powers; Algebra",
             "acs": "AM.I.H.K1 – AM.I.H.K4",
             "quiz": "Quiz 1 from Fundamentals of Math - Day 1"},
        ],
    )
    ws2 = wb.create_sheet("Block 2 (Day-Night)")
    _write_sheet(
        ws2,
        banner="Block 2 (Day/Night): Aircraft Drawings, Materials and Processes",
        subject="Aircraft Drawings",
        days=[
            {"day": 1, "subj_day": 1, "topics": "Introduction; Title Blocks",
             "acs": "AM.I.B.K1, K2, K4, R1, R3",
             "handbook": "FAA-H-8083-30B • Ch. 4 pp. 4-1 to 4-9"},
            {"day": 2, "subj_day": 2, "topics": "Drawing Symbols",
             "acs": "AM.I.B.K1 – AM.I.B.K4",
             "projects": "Project 2-1",
             "quiz": "Quiz 1 from Aircraft Drawings - Day 1"},
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _legacy_single_sheet_bytes() -> bytes:
    """Old layout: Subject/Day already says 'Aircraft Drawings - Day 1'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Block 6 Teacher Calendar"
    headers = ["Days", "Subject/Day", "Topics Covered", "ACS Codes", "Handbook",
               "Projects", "ACS Codes for Projects", "Quiz"]
    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h)
    ws.cell(2, 1, 1)
    ws.cell(2, 2, "Aircraft Drawings - Day 1")
    ws.cell(2, 3, "Orthographic projection")
    ws.cell(2, 4, "AM.I.B.K1")
    ws.cell(2, 5, "FAA-H-8083-30B Ch. 2 pgs 2-8 to 2-14")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Parser unit tests
# ---------------------------------------------------------------------------

def test_expand_acs_shorthand_after_full_code():
    assert expand_acs("AM.I.B.K1, K2, K4, R1, R3") == [
        "AM.I.B.K1", "AM.I.B.K2", "AM.I.B.K4", "AM.I.B.R1", "AM.I.B.R3",
    ]


def test_parse_handbook_accepts_pp():
    refs = parse_handbook("FAA-H-8083-30B • Ch. 4 pp. 4-1 to 4-9")
    assert len(refs) == 1
    assert refs[0]["handbook"] == "FAA-H-8083-30B"
    assert refs[0]["chapter"] == 4
    assert "4-1" in refs[0]["pages"]


def test_multi_sheet_workbook_parses_all_sheets_and_ignores_download_suffix():
    content = _two_sheet_workbook_bytes()
    assert looks_like_aim_teacher_calendar(content)
    # Filename suffix "(1)" must NOT become Block 1 for the whole book.
    result = build_calendar_structure(
        content, "ALL Block Calendars_NEW FORMAT (1).xlsx",
    )
    assert len(result["sheets"]) == 2
    assert result["days"] == [], "multi-sheet must leave top-level days empty"
    assert result["block"] == ""
    names = [s["sheet_name"] for s in result["sheets"]]
    assert names == ["Block 1 (Day-Night)", "Block 2 (Day-Night)"]
    assert result["sheets"][0]["block"] == "Block 1"
    assert result["sheets"][1]["block"] == "Block 2"
    assert result["sheets"][0]["schedule"] == "day_night"
    assert result["blocks_covered"] == ["Block 1", "Block 2"]


def test_new_format_subject_header_and_acs_and_handbook():
    content = _two_sheet_workbook_bytes()
    result = build_calendar_structure(content, "calendars.xlsx")
    b2 = result["sheets"][1]
    day1 = b2["days"][0]
    assert day1["subject_unit"] == "Aircraft Drawings"
    assert day1["subject_day"] == 1
    assert "Aircraft Drawings" in day1["lesson_title"]
    assert "AM.I.B.K2" in day1["acs_codes"]
    assert "AM.I.B.R1" in day1["acs_codes"]
    assert day1["handbook_refs"][0]["chapter"] == 4
    day2 = b2["days"][1]
    assert day2["quiz"] and day2["quiz"]["number"] == 1
    assert day2["quiz"]["covers_days"] == [1]


def test_legacy_single_sheet_still_fills_top_level_days():
    content = _legacy_single_sheet_bytes()
    result = build_calendar_structure(content, "Block 6 Teacher Calendar.xlsx")
    assert len(result["sheets"]) == 1
    assert result["block"] == "Block 6"
    assert len(result["days"]) == 1
    assert result["days"][0]["subject_unit"] == "Aircraft Drawings"
    assert result["days"][0]["handbook_refs"][0]["chapter"] == 2


# ---------------------------------------------------------------------------
# upsert_calendar
# ---------------------------------------------------------------------------

class _Cur:
    """Records INSERTs into calendar tables; stands in for a psycopg cursor."""

    def __init__(self):
        self.calendar_rows = []
        self.day_rows = []
        self.deletes = []

    def execute(self, sql, params=None):
        sql_l = sql.lower()
        if sql_l.strip().startswith("delete"):
            self.deletes.append((sql, params))
            return
        if "dis_calendar_days" in sql_l and "insert" in sql_l:
            self.day_rows.append(params)
        if "dis_course_calendars" in sql_l and "insert" in sql_l:
            self.calendar_rows.append(params)


def _state(day_numbers):
    return {
        "job_id": "job1", "tenant_id": "aim", "client_id": "aim",
        "calendar_structure": {
            "block": "Block 13", "course_name": "AIM",
            "days": [{"day_number": n, "topic": f"t{i}"} for i, n in enumerate(day_numbers)],
        },
    }


def test_a_calendar_that_collapses_says_so():
    state = _state([1] * 11)
    stored = upsert_calendar(_Cur(), "dis", "doc1", state, "test")
    assert stored == 1
    assert state["errors"]
    assert "day_number" in state["errors"][0] and "overwritten" in state["errors"][0]


def test_the_returned_count_is_days_actually_stored():
    state = _state([1, 1, 2, 2, 3])
    assert upsert_calendar(_Cur(), "dis", "doc1", state, "test") == 3


def test_a_clean_calendar_reports_nothing():
    state = _state(list(range(1, 21)))
    assert upsert_calendar(_Cur(), "dis", "doc1", state, "test") == 20
    assert not state.get("errors")


def test_day_rows_carry_the_canonical_block_label():
    state = _state([1, 2])
    state["calendar_structure"]["block"] = "Block 09"
    cur = _Cur()
    upsert_calendar(cur, "dis", "doc1", state, "test")
    params = [p for p in cur.day_rows[0] if isinstance(p, str)]
    assert "Block 9" in params and "Block 09" not in params
    assert "test" in params


def test_two_sheets_with_same_day_number_do_not_collide():
    state = {
        "job_id": "job_multi", "tenant_id": "aim", "client_id": "aim",
        "calendar_structure": {
            "course_name": "AIM",
            "sheets": [
                {
                    "sheet_name": "Block 2 (Day-Night)", "sheet_index": 0,
                    "block": "Block 2", "block_number": 2, "schedule": "day_night",
                    "days": [{"day_number": 1, "topic": "DN day 1", "lesson_title": "DN"}],
                },
                {
                    "sheet_name": "Block 2 (Weekend)", "sheet_index": 1,
                    "block": "Block 2", "block_number": 2, "schedule": "weekend",
                    "days": [{"day_number": 1, "topic": "WE day 1", "lesson_title": "WE"}],
                },
            ],
            "days": [],
        },
    }
    cur = _Cur()
    stored = upsert_calendar(cur, "dis", "doc_multi", state, "test")
    assert stored == 2
    assert len(cur.calendar_rows) == 2
    assert len(cur.day_rows) == 2
    cal_ids = {r[0] for r in cur.calendar_rows}
    assert cal_ids == {"cal_job_multi:s0", "cal_job_multi:s1"}
    day_ids = {r[0] for r in cur.day_rows}
    assert day_ids == {"cal_job_multi:s0:day_1", "cal_job_multi:s1:day_1"}
    assert len(cur.deletes) == 2


# ---------------------------------------------------------------------------
# Source Library per_unit
# ---------------------------------------------------------------------------

def test_course_calendar_is_per_unit_not_full_document():
    assert "course_calendar" in PER_UNIT_DOC_TYPES
    assert chunking_strategy("blueprint", "course_calendar") == "per_unit"


def test_build_clean_content_keeps_sheet_titles():
    payload = {
        "job_id": "jobx",
        "tenant_id": "aim",
        "client_id": "aim",
        "metadata": {
            "document_type": "course_calendar",
            "purpose": "blueprint",
            "title": "ALL Block Calendars",
        },
        "source_file": {"name": "ALL Block Calendars_NEW FORMAT.xlsx", "type": "xlsx"},
        "content_units": [
            {
                "content_unit_id": "jobx:calendar_sheet_0",
                "unit_type": "calendar_sheet",
                "unit_number": 1,
                "title": "Block 1 (Day-Night)",
                "text": "Block 1 day text",
                "metadata": {
                    "block": "Block 1", "sheet_name": "Block 1 (Day-Night)",
                    "sheet_index": 0, "schedule": "day_night", "total_days": 2,
                    "acs_codes": ["AM.I.H.K3"],
                },
            },
            {
                "content_unit_id": "jobx:calendar_sheet_1",
                "unit_type": "calendar_sheet",
                "unit_number": 2,
                "title": "Block 2 (Day-Night)",
                "text": "Block 2 day text",
                "metadata": {
                    "block": "Block 2", "sheet_name": "Block 2 (Day-Night)",
                    "sheet_index": 1, "schedule": "day_night", "total_days": 2,
                    "acs_codes": ["AM.I.B.K1"],
                },
            },
        ],
    }
    doc = build_clean_content_document(payload)
    assert doc["chunking_strategy"] == "per_unit"
    assert doc["total_units"] == 2
    titles = [u["title"] for u in doc["content_units"]]
    assert titles == ["Block 1 (Day-Night)", "Block 2 (Day-Night)"]
    assert doc["content_units"][1]["metadata"]["block"] == "Block 2"
    assert doc["content_units"][1]["metadata"]["sheet_name"] == "Block 2 (Day-Night)"


def test_iter_calendar_sheets_wraps_legacy_days():
    sheets = iter_calendar_sheets({
        "block": "Block 6", "block_number": 6,
        "days": [{"day_number": 1, "topic": "t"}],
    })
    assert len(sheets) == 1
    assert sheets[0]["block"] == "Block 6"
    assert len(sheets[0]["days"]) == 1
