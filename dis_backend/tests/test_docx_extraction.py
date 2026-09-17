"""Final-exam DOCX files were ingested as empty because python-docx only reads
body paragraphs and tables.

AIM's "POC Block 05 Final Cumulative Exam" / "Figures 1–7" uploads put questions
in text boxes, DrawingML, image alt text, or a pre-2007 OLE2 container saved
with a .docx extension. Those all produced extracted_chars=0, Source Library
"Nothing extracted", and nothing in the index. These tests pin the XML walk and
the OLE scrape that recover that text without antiword.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document

from services.pipeline.extractors import extract, extract_docx


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_OLE2_MAGIC_PAD = b"\xd0\xcf\x11\xe0" + b"\x00" * 64


def _docx_zip(**parts: str | bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="png" ContentType="image/png"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        for name, payload in parts.items():
            zf.writestr(name.replace("__", "/"), payload)
    return buf.getvalue()


def _python_docx_bytes(paragraphs) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_regular_docx_body_text_still_extracts():
    content = _python_docx_bytes([
        "Block 05 Final Cumulative Exam",
        "Question 1: Identify the metallic structure.",
    ])
    result = extract("POC_Block_05_Final_Cumulative_Exam_updated.docx", content)
    assert "Question 1: Identify the metallic structure." in result.text
    assert "Final Cumulative Exam" in result.text


def test_textbox_only_docx_is_not_empty():
    """python-docx ignores w:txbxContent; that is the empty-exam failure mode."""
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body>
    <w:p>
      <w:r>
        <w:drawing>
          <w:txbxContent>
            <w:p><w:r><w:t>Question 1: Identify the landing gear component shown.</w:t></w:r></w:p>
            <w:p><w:r><w:t>A. Nose strut  B. Torque link  C. Shimmy damper</w:t></w:r></w:p>
          </w:txbxContent>
        </w:drawing>
      </w:r>
    </w:p>
  </w:body>
</w:document>"""
    content = _docx_zip(**{"word/document.xml": document_xml})

    # The original python-docx paragraph walk sees nothing here (or cannot
    # open the sparse zip at all). Either way the exam used to index empty.
    try:
        doc = Document(io.BytesIO(content))
        assert not any(p.text.strip() for p in doc.paragraphs)
    except Exception:
        pass

    result = extract_docx(content)
    assert "Question 1: Identify the landing gear component shown." in result.text
    assert "Shimmy damper" in result.text


def test_drawingml_and_image_alt_text_are_indexed():
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}" xmlns:a="{A_NS}" xmlns:wp="{WP_NS}">
  <w:body>
    <w:p>
      <w:r>
        <w:drawing>
          <wp:docPr id="1" name="Figure 1" descr="Nose wheel strut with labeled torque links"/>
          <a:p><a:r><a:t>Figure 1 — Nose wheel assembly</a:t></a:r></a:p>
        </w:drawing>
      </w:r>
    </w:p>
  </w:body>
</w:document>"""
    content = _docx_zip(
        **{
            "word/document.xml": document_xml,
            "word/media/image1.png": b"\x89PNG\r\n\x1a\n" + b"\x00" * 16,
        }
    )
    result = extract_docx(content)
    assert "Figure 1 — Nose wheel assembly" in result.text
    assert "Nose wheel strut with labeled torque links" in result.text
    assert result.has_images is True


def test_header_text_is_included():
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}">
  <w:body><w:p><w:r><w:t>Question 4: Select the correct alloy.</w:t></w:r></w:p></w:body>
</w:document>"""
    header_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr xmlns:w="{W_NS}">
  <w:p><w:r><w:t>Block 05 Final Cumulative Exam — Student Copy</w:t></w:r></w:p>
</w:hdr>"""
    content = _docx_zip(**{
        "word/document.xml": document_xml,
        "word/header1.xml": header_xml,
    })
    result = extract_docx(content)
    assert "Question 4" in result.text
    assert "Student Copy" in result.text


def test_legacy_ole2_docx_scrapes_utf16_when_antiword_is_missing(monkeypatch):
    question = "Question 1: What is the purpose of a torque link?"
    # OLE2 magic + UTF-16LE body. Not a real Word FIB; the scrape path must
    # still recover the exam text that antiword would have returned.
    payload = _OLE2_MAGIC_PAD + question.encode("utf-16le")
    monkeypatch.setattr("shutil.which", lambda name: None)
    result = extract_docx(payload)
    assert "torque link" in result.text.lower()


def test_ocr_is_not_invoked_on_the_fast_upload_path(monkeypatch):
    called = {"n": 0}

    def boom(_content):
        called["n"] += 1
        raise AssertionError("OCR must not run on the upload preview path")

    monkeypatch.setattr("services.pipeline.extractors._ocr_ooxml_images", boom)
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body><w:p/></w:body></w:document>"""
    content = _docx_zip(**{
        "word/document.xml": document_xml,
        "word/media/image1.png": b"\x89PNG\r\n\x1a\n" + b"\x00" * 16,
    })
    result = extract_docx(content)  # options default: no OCR
    assert called["n"] == 0
    assert result.has_images is True
    assert not result.text.strip()


def test_final_exam_content_type_gets_quiz_structure():
    """AIM stores exams as content_type=final_exam after classification.

    Specialized structure used to key only on pipeline doc_type=quiz_exam, so a
    remapped exam (or one the LLM left as `other`) never got quiz_structure.
    """
    from types import SimpleNamespace

    from config.settings import DocumentProcessingConfig
    from services.agents.specialized_structure_extraction_agent import (
        SpecializedStructureExtractionAgent,
    )

    class _Ctx:
        def __init__(self):
            self.cfg = SimpleNamespace(document_processing=DocumentProcessingConfig())

        def step_done(self, state, name):
            return state

    agent = SpecializedStructureExtractionAgent(_Ctx())
    state = {
        "doc_type": "other",
        "doc_metadata": {"content_type": "final_exam"},
        "filename": "POC_Block_05_Final_Cumulative_Exam_updated.docx",
        "raw_text": "Question 1: What is a torque link?\nQuestion 2: Name the alloy.",
        "tables": [],
    }
    out = agent.run(state)
    assert out["quiz_structure"]["question_count_detected"] == 2
    assert out["specialized_structure_type"]
