"""Style page renderer.

Extracted from the legacy Streamlit monolith so each page can be maintained independently.
"""

import re
import streamlit as st
from datetime import datetime, timezone
from concurrent.futures import ThreadPoolExecutor

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

from promptops_app.database import (
    Document, Style,
    create_style, add_files_to_style, get_styles,
    get_active_style, set_active_style, deactivate_style,
    log_event,
)
from promptops_app.auth.permissions import rbac_gate, can_modify_style
from promptops_app.parsers.file_parser import _parse_uploaded_file
from promptops_app.services.style_service import (
    generate_style_understanding, regenerate_style_understanding,
)
from promptops_app.core.config import DOCUMENT_PAGE_SIZE
from promptops_app.repositories import document_repository
from promptops_app.services.audit_service import log_audit_event
from promptops_app.ui.components import _section_badge
from promptops_app.ui.pagination import paginate as _paginate_docs
from promptops_app.ui.prompt_panel import render_prompt_panel
from promptops_app.ui.generation_controls import render_inline_prompt_controls, render_prompt_download_button
from promptops_app.ui.user_prompt_widget import auto_save_instructions
from promptops_app.prompt_templates import CDD_SYSTEM_PROMPT


def render_page(db, ctx):
    user_role = ctx.user_role
    user_name = ctx.user_name
    _proj_id = ctx.project_id
    _proj_name = ctx.project_name
    _crs_id = ctx.course_id
    _crs_name = ctx.course_name
    _is_admin = ctx.is_admin
    _is_lead = ctx.is_lead
    model_choice = ctx.model_choice
    expert_domain = ctx.expert_domain
    target_audience = ctx.target_audience
    aud_cat = ctx.audience_category
    # =====================================================================
    # STYLE — Style Management + Document Registry (two sub-tabs)
    # =====================================================================
    st.markdown(_section_badge("🎨", "Style Management & Document Registry",
        "Create instructional Styles that govern tone and structure across every generation, "
        "and manage the Context Document Database."),
        unsafe_allow_html=True)

    # ── Active Style Banner (scoped to current course/project) ───────────
    _active_style = get_active_style(db, project_id=_proj_id, course_id=_crs_id)
    if _active_style:
        st.markdown(
            f"<div style='background:#f0fdf4;border:1.5px solid #86efac;border-radius:10px;"
            f"padding:10px 16px;margin-bottom:12px;display:flex;align-items:center;gap:10px;'>"
            f"<span style='font-size:1.2rem;'>🎨</span>"
            f"<div><span style='font-size:0.7rem;font-weight:700;text-transform:uppercase;"
            f"letter-spacing:.08em;color:#166534;'>Active Style</span><br>"
            f"<strong style='color:#15803d;font-size:0.95rem;'>{_active_style.name}</strong>"
            f"<span style='font-size:0.78rem;color:#6b7280;margin-left:10px;'>"
            f"Auto-injected into CDD → Blueprint → Generate</span></div></div>",
            unsafe_allow_html=True
        )
    else:
        st.markdown(
            "<div style='background:#fef9c3;border:1px solid #fde047;border-radius:8px;"
            "padding:8px 14px;font-size:0.82rem;color:#854d0e;margin-bottom:12px;'>"
            "⚠️ <strong>No active style.</strong> Create and activate a Style below for consistent "
            "tone and structure across all generations.</div>",
            unsafe_allow_html=True
        )

    _can_modify_style = can_modify_style(db, user_name, user_role, project_id=_proj_id)

    if not _can_modify_style:
        st.markdown(
            "<div style='background:#f0f9ff;border:1.5px solid #7dd3fc;border-radius:10px;"
            "padding:10px 16px;margin-bottom:12px;display:flex;align-items:center;gap:10px;'>"
            "<span style='font-size:1.2rem;'>🔍</span>"
            "<div><span style='font-size:0.7rem;font-weight:700;text-transform:uppercase;"
            "letter-spacing:.08em;color:#0369a1;'>Read-Only Access</span><br>"
            "<span style='font-size:0.85rem;color:#0c4a6e;'>Style configuration is managed by Admin and Lead. "
            "Your workflow starts from the <strong>CDD</strong> tab.</span></div></div>",
            unsafe_allow_html=True
        )

    ctx_tab1, ctx_tab2 = st.tabs(["🎨 Style Management", "📄 Document Registry"])

    # =====================================================================
    # TAB 1 — STYLE MANAGEMENT
    # =====================================================================
    with ctx_tab1:
        sty_left, sty_right = st.columns([0.42, 0.58], gap="large")

        # ── LEFT PANEL: Create / Edit Style ──────────────────────────────
        with sty_left:
            if not _can_modify_style:
                st.markdown(
                    "<div style='background:#f8fafc;border:1px solid #e2e8f0;border-radius:10px;"
                    "padding:20px 18px;text-align:center;'>"
                    "<div style='font-size:2rem;margin-bottom:8px;'>🔒</div>"
                    "<div style='font-weight:700;color:#374151;margin-bottom:4px;'>Style Creation Restricted</div>"
                    "<div style='font-size:0.82rem;color:#6b7280;'>Only Admins and Leads can create or "
                    "modify Styles. Contact your Lead or Admin to update Style configuration.</div>"
                    "</div>",
                    unsafe_allow_html=True
                )
            else:
                st.markdown("#### ➕ Create New Style")
                st.markdown(
                    "<p style='font-size:0.82rem;color:#6b7280;margin-top:-0.4rem;margin-bottom:0.8rem;'>"
                    "A Style defines writing tone, rules, and structure. It will be automatically "
                    "applied to CDD, Blueprint, and content generation.</p>",
                    unsafe_allow_html=True
                )

            if _can_modify_style:
                all_active_docs_sty = document_repository.list_active_documents(db)

                with st.form("style_create_form"):
                    sty_name = st.text_input(
                        "Style Name *",
                        placeholder="e.g. Clinical Nursing — Formal Academic"
                    )
                    sty_desc = st.text_input(
                        "Description",
                        placeholder="Brief description of this style's purpose"
                    )

                    st.markdown(
                        "<div style='font-size:0.75rem;font-weight:700;text-transform:uppercase;"
                        "letter-spacing:.08em;color:#4f46e5;margin:8px 0 4px;'>📎 Reference Documents</div>",
                        unsafe_allow_html=True
                    )
                    # Option A: select from existing library (multi-select)
                    _lib_filenames = [d.filename for d in all_active_docs_sty]
                    sty_sel_docs = st.multiselect(
                        "Select existing documents from library",
                        _lib_filenames,
                        help="These documents define tone, guidelines, and structure for this style."
                    )
                    # Option B: upload multiple new files simultaneously
                    sty_upload_files = st.file_uploader(
                        "Upload new style reference files (multiple allowed)",
                        type=["pdf", "docx", "txt"],
                        accept_multiple_files=True,
                        key="sty_new_files"
                    )

                    st.markdown(
                        "<div style='font-size:0.75rem;font-weight:700;text-transform:uppercase;"
                        "letter-spacing:.08em;color:#4f46e5;margin:8px 0 4px;'>✍️ Custom Instructions</div>",
                        unsafe_allow_html=True
                    )
                    sty_instructions = st.text_area(
                        "Custom instructions / rules for this style",
                        placeholder=(
                            "e.g. Always use active voice. "
                            "Use Bloom's taxonomy levels 3–5. "
                            "Avoid jargon. "
                            "Begin each lesson with a real-world clinical scenario."
                        ),
                        height=130,
                        help="These instructions are stored with the style and injected into every generation."
                    )

                    _sty_b1, _sty_b2, _sty_b3 = st.columns([0.25, 0.5, 0.25])
                    sty_create_btn = _sty_b2.form_submit_button(
                        "💾 Save Style", use_container_width=True, type="primary"
                    )

                if sty_create_btn:
                    if not rbac_gate(user_role, "style.create", "Creating a Style"):
                        pass  # rbac_gate already shows the error
                    elif not sty_name.strip():
                        st.warning("Style Name is required.")
                    else:
                        # Parse all uploaded files in parallel, then write to DB sequentially
                        _new_doc_ids = []
                        if sty_upload_files:
                            with ThreadPoolExecutor(max_workers=min(4, len(sty_upload_files))) as _pool:
                                _parsed = list(_pool.map(_parse_uploaded_file, sty_upload_files))
                            _file_type_map = {uf.name: uf.type for uf in sty_upload_files}
                            for _fname, _uf_content, _err in _parsed:
                                if _err:
                                    st.error(f"❌ Error reading {_fname}: {_err}")
                                    continue
                                if _uf_content:
                                    _now_sty = datetime.now(timezone.utc)
                                    _existing_sty_doc = document_repository.get_document_by_filename(db, _fname)
                                    if _existing_sty_doc:
                                        _new_doc_ids.append(_existing_sty_doc.id)
                                    else:
                                        _new_doc = Document(
                                            filename=_fname,
                                            file_type=_file_type_map.get(_fname, ""),
                                            doc_tag="style_reference",
                                            content=_uf_content,
                                            uploaded_by=user_name,
                                            uploaded_at=_now_sty,
                                            updated_at=_now_sty,
                                            status="active"
                                        )
                                        db.add(_new_doc); db.commit(); db.refresh(_new_doc)
                                        _new_doc_ids.append(_new_doc.id)

                        # Gather doc IDs from library selection
                        _sel_doc_ids = [
                            d.id for d in all_active_docs_sty if d.filename in sty_sel_docs
                        ]
                        _all_doc_ids = list(set(_sel_doc_ids + _new_doc_ids))

                        new_style = create_style(
                            db, sty_name.strip(), sty_desc.strip(),
                            sty_instructions.strip(), _all_doc_ids, user_name
                        )
                        # Auto-activate for current course scope
                        set_active_style(db, new_style.id, scope="course", course_id=_crs_id)
                        log_event(db, "style_created", user_name,
                            f"Style '{new_style.name}' created & activated for course {_crs_id}",
                            {"style_id": new_style.style_id, "doc_count": len(_all_doc_ids)})
                        log_audit_event(db, user_name, "style.uploaded", entity_type="style", entity_id=new_style.style_id,
                                        project_id=_proj_id, course_id=_crs_id,
                                        metadata={"name": new_style.name, "doc_count": len(_all_doc_ids)})
                        st.toast(f"✅ Style '{new_style.name}' created and activated for this course.")
                        st.rerun()

        # ── RIGHT PANEL: Saved Styles List & Management ───────────────────
        with sty_right:
            st.markdown("#### 🗂️ Saved Styles")
            all_styles = get_styles(db)

            if not all_styles:
                st.info("No styles yet. Create your first style using the form on the left.")
            else:
                for sty in all_styles:
                    _sty_is_active = sty.is_active
                    _sty_bg = "#f0fdf4" if _sty_is_active else "#ffffff"
                    _sty_border = "#86efac" if _sty_is_active else "#e4e7ef"
                    _sty_docs = [sd.document for sd in sty.style_documents if sd.document]

                    with st.container():
                        st.markdown(
                            f"<div style='background:{_sty_bg};border:1.5px solid {_sty_border};"
                            f"border-radius:10px;padding:12px 16px;margin-bottom:4px;'>"
                            f"<div style='display:flex;align-items:center;gap:8px;'>"
                            f"<span style='font-size:1rem;'>{'✅' if _sty_is_active else '🎨'}</span>"
                            f"<strong style='font-size:0.95rem;color:#111827;'>{sty.name}</strong>"
                            + (f"<span style='background:#dcfce7;color:#166534;font-size:0.68rem;"
                               f"font-weight:700;padding:2px 7px;border-radius:10px;margin-left:6px;'>"
                               f"ACTIVE</span>" if _sty_is_active else "")
                            + f"</div>"
                            f"<div style='font-size:0.78rem;color:#6b7280;margin-top:4px;'>"
                            f"{sty.description or '—'}</div>"
                            f"<div style='font-size:0.72rem;color:#9ca3af;margin-top:3px;'>"
                            f"📎 {len(_sty_docs)} doc(s) · Created {sty.created_at.strftime('%Y-%m-%d') if sty.created_at else '—'}"
                            f"</div></div>",
                            unsafe_allow_html=True
                        )

                        # Action row — View is always visible; modification controls restricted to Admin/Lead
                        if _can_modify_style:
                            _sa1, _sa2, _sa3, _sa4, _sa5, _sa6 = st.columns([0.16, 0.17, 0.17, 0.17, 0.17, 0.16])
                        else:
                            _sa1 = st.columns([0.25, 0.75])[0]

                        # View details toggle — all roles
                        if _sa1.button("👁️ View", key=f"sty_view_{sty.id}", use_container_width=True):
                            st.session_state[f"sty_open_{sty.id}"] = not st.session_state.get(f"sty_open_{sty.id}", False)

                        if _can_modify_style:
                            # Generate Understanding
                            if _sa2.button("🧠 Understand", key=f"sty_gen_{sty.id}", use_container_width=True):
                                st.session_state[f"sty_gen_open_{sty.id}"] = not st.session_state.get(f"sty_gen_open_{sty.id}", False)

                            # Regenerate / Refine
                            if _sa3.button("🔄 Refine", key=f"sty_refine_{sty.id}", use_container_width=True):
                                st.session_state[f"sty_refine_open_{sty.id}"] = not st.session_state.get(f"sty_refine_open_{sty.id}", False)

                            # Add Files to existing style
                            if _sa4.button("➕ Files", key=f"sty_addfiles_btn_{sty.id}", use_container_width=True):
                                st.session_state[f"sty_addfiles_open_{sty.id}"] = not st.session_state.get(f"sty_addfiles_open_{sty.id}", False)

                            # Set as Active — scoped
                            _sty_scope_key = f"sty_scope_{sty.id}"
                            _sty_scope_choice = st.session_state.get(_sty_scope_key, "Course")
                            if not _sty_is_active:
                                if _sa5.button("📌 Activate", key=f"sty_activate_{sty.id}", use_container_width=True, type="primary"):
                                    if rbac_gate(user_role, "style.activate", "Activating a Style"):
                                        st.session_state[f"sty_scope_open_{sty.id}"] = True
                                        st.rerun()
                            else:
                                if _sa5.button("⏹ Deactivate", key=f"sty_deact_{sty.id}", use_container_width=True):
                                    if rbac_gate(user_role, "style.activate", "Deactivating a Style"):
                                        deactivate_style(db, scope="course", course_id=_crs_id)
                                        deactivate_style(db, scope="project", project_id=_proj_id)
                                        if _is_admin:
                                            deactivate_style(db, scope="global")
                                        log_event(db, "style_deactivated", user_name, f"Style '{sty.name}' deactivated")
                                        st.toast(f"Style '{sty.name}' deactivated.")
                                        st.rerun()

                            # Scope picker popup
                            if st.session_state.get(f"sty_scope_open_{sty.id}"):
                                with st.container():
                                    st.markdown(
                                        "<div style='background:#f0f9ff;border:1px solid #bae6fd;"
                                        "border-radius:8px;padding:10px 14px;margin:6px 0;'>"
                                        "<div style='font-size:0.82rem;font-weight:700;color:#0369a1;"
                                        "margin-bottom:6px;'>📌 Choose activation scope</div></div>",
                                        unsafe_allow_html=True
                                    )
                                    _sc_cols = st.columns(3)
                                    if _sc_cols[0].button("For This Course", key=f"sty_act_crs_{sty.id}", use_container_width=True, type="primary"):
                                        set_active_style(db, sty.id, scope="course", course_id=_crs_id)
                                        log_event(db, "style_activated", user_name, f"Style '{sty.name}' activated for course {_crs_id}")
                                        st.session_state.pop(f"sty_scope_open_{sty.id}", None)
                                        st.toast(f"✅ '{sty.name}' active for this course.")
                                        st.rerun()
                                    if _sc_cols[1].button("For This Project", key=f"sty_act_proj_{sty.id}", use_container_width=True):
                                        set_active_style(db, sty.id, scope="project", project_id=_proj_id)
                                        log_event(db, "style_activated", user_name, f"Style '{sty.name}' activated for project {_proj_id}")
                                        st.session_state.pop(f"sty_scope_open_{sty.id}", None)
                                        st.toast(f"✅ '{sty.name}' active for this project.")
                                        st.rerun()
                                    if _is_admin and _sc_cols[2].button("Globally", key=f"sty_act_global_{sty.id}", use_container_width=True):
                                        set_active_style(db, sty.id, scope="global")
                                        log_event(db, "style_activated", user_name, f"Style '{sty.name}' activated globally")
                                        st.session_state.pop(f"sty_scope_open_{sty.id}", None)
                                        st.toast(f"✅ '{sty.name}' active globally.")
                                        st.rerun()
                                    if st.button("Cancel", key=f"sty_scope_cancel_{sty.id}"):
                                        st.session_state.pop(f"sty_scope_open_{sty.id}", None)
                                        st.rerun()

                            # Delete
                            if _sa6.button("🗑️ Delete", key=f"sty_del_{sty.id}", use_container_width=True):
                                if rbac_gate(user_role, "style.delete", "Deleting a Style"):
                                    db.delete(sty); db.commit()
                                    log_event(db, "style_deleted", user_name, f"Style '{sty.name}' deleted")
                                    st.toast(f"🗑️ '{sty.name}' deleted.")
                                    st.rerun()

                        # ── Stale Understanding Banner ────────────────────
                        _und_status = getattr(sty, "understanding_status", "fresh") or "fresh"
                        if _und_status == "stale":
                            st.markdown(
                                "<div style='background:#fff7ed;border:1.5px solid #fb923c;"
                                "border-radius:8px;padding:10px 14px;margin:8px 0 4px;display:flex;"
                                "align-items:center;gap:10px;'>"
                                "<span style='font-size:1.1rem;'>⚠️</span>"
                                "<div style='flex:1;font-size:0.83rem;color:#92400e;'>"
                                "<strong>Understanding is out of date.</strong> New files were added. "
                                "An Admin or Lead needs to regenerate the understanding.</div>"
                                "</div>",
                                unsafe_allow_html=True
                            )
                            if _can_modify_style and st.button(
                                "🔄 Regenerate Understanding",
                                key=f"sty_stale_regen_{sty.id}",
                                type="primary",
                                use_container_width=False
                            ):
                                st.session_state[f"sty_gen_open_{sty.id}"] = True
                                st.rerun()

                        # ── Add Files Panel — Admin/Lead only ─────────────
                        if st.session_state.get(f"sty_addfiles_open_{sty.id}", False) and _can_modify_style:
                            with st.expander(f"➕ Add Files to Style — {sty.name}", expanded=True):
                                st.markdown(
                                    "<p style='font-size:0.82rem;color:#6b7280;'>"
                                    "Upload additional reference files. They will be appended to the existing "
                                    "file list — no existing files will be removed.</p>",
                                    unsafe_allow_html=True
                                )
                                _all_lib_docs = document_repository.list_active_documents(db)
                                _existing_doc_ids = {sd.document_id for sd in sty.style_documents}
                                _available_lib = [d for d in _all_lib_docs if d.id not in _existing_doc_ids]

                                _af_lib = st.multiselect(
                                    "Add from document library",
                                    [d.filename for d in _available_lib],
                                    help="Only documents not already linked to this style are shown.",
                                    key=f"sty_af_lib_{sty.id}"
                                )
                                _af_upload = st.file_uploader(
                                    "Upload new reference files",
                                    type=["pdf", "docx", "txt"],
                                    accept_multiple_files=True,
                                    key=f"sty_af_upload_{sty.id}"
                                )
                                _af_instructions = st.text_input(
                                    "Additional Instructions (optional)",
                                    placeholder="e.g. Focus more on assessment tone",
                                    key=f"sty_af_instr_{sty.id}"
                                )

                                if st.button("📎 Append Files", key=f"sty_af_commit_{sty.id}", type="primary"):
                                    if rbac_gate(user_role, "style.upload", "Uploading files to a Style"):
                                        _af_new_ids = []
                                        if _af_upload:
                                            with ThreadPoolExecutor(max_workers=min(4, len(_af_upload))) as _af_pool:
                                                _af_parsed = list(_af_pool.map(_parse_uploaded_file, _af_upload))
                                            _af_type_map = {uf.name: uf.type for uf in _af_upload}
                                            for _af_fname, _af_content, _af_err in _af_parsed:
                                                if _af_err:
                                                    st.error(f"❌ Error reading {_af_fname}: {_af_err}")
                                                    continue
                                                if _af_content:
                                                    _af_now = datetime.now(timezone.utc)
                                                    _af_existing = document_repository.get_document_by_filename(db, _af_fname)
                                                    if _af_existing:
                                                        _af_new_ids.append(_af_existing.id)
                                                    else:
                                                        _af_doc = Document(
                                                            filename=_af_fname,
                                                            file_type=_af_type_map.get(_af_fname, ""),
                                                            doc_tag="style_reference", content=_af_content,
                                                            uploaded_by=user_name, uploaded_at=_af_now,
                                                            updated_at=_af_now, status="active"
                                                        )
                                                        db.add(_af_doc); db.commit(); db.refresh(_af_doc)
                                                        _af_new_ids.append(_af_doc.id)

                                        _af_lib_ids = [d.id for d in _available_lib if d.filename in _af_lib]
                                        _af_all_new = list(set(_af_lib_ids + _af_new_ids))

                                        if not _af_all_new:
                                            st.warning("No new files selected or uploaded.")
                                        else:
                                            _af_added = add_files_to_style(db, sty, _af_all_new)
                                            if _af_instructions.strip():
                                                st.session_state[f"sty_pending_instructions_{sty.id}"] = _af_instructions.strip()
                                            log_event(db, "style_files_added", user_name,
                                                f"{_af_added} file(s) appended to style '{sty.name}'",
                                                {"style_id": sty.style_id, "added": _af_added})
                                            st.toast(f"✅ {_af_added} file(s) added. Understanding marked as stale.")
                                            st.session_state[f"sty_addfiles_open_{sty.id}"] = False
                                            st.rerun()

                        # ── View Panel ────────────────────────────────────
                        if st.session_state.get(f"sty_open_{sty.id}", False):
                            with st.expander(f"📋 Style Details — {sty.name}", expanded=True):
                                st.markdown(f"**Style ID:** `{sty.style_id}`")
                                if sty.custom_instructions:
                                    st.markdown("**Custom Instructions:**")
                                    st.text_area("", value=sty.custom_instructions, height=100,
                                                 disabled=True, key=f"sty_inst_{sty.id}")
                                if _sty_docs:
                                    st.markdown(f"**Reference Documents ({len(_sty_docs)}):**")
                                    for _d in _sty_docs:
                                        st.markdown(f"- 📄 `{_d.filename}` ({_d.doc_tag or 'general'})")
                                if sty.generated_summary:
                                    st.markdown("---")
                                    st.markdown("**🧠 Style Intelligence Layer (stored understanding):**")
                                    # Render the 4-section format as distinct blocks
                                    _und_text = sty.generated_summary
                                    _und_sections = {
                                        "WHAT THIS IS": ("#eef2ff", "#4338ca", "📌"),
                                        "WHAT I LEARNED": ("#f0fdf4", "#166534", "📖"),
                                        "HOW I WILL WORK": ("#fff7ed", "#92400e", "⚙️"),
                                        "WHAT I WILL NOT DO": ("#fef2f2", "#991b1b", "🚫"),
                                    }
                                    import re as _re_sty
                                    for _sec_name, (_sbg, _sfg, _sico) in _und_sections.items():
                                        _pat = rf"{re.escape(_sec_name)}\s*(.*?)(?=WHAT THIS IS|WHAT I LEARNED|HOW I WILL WORK|WHAT I WILL NOT DO|$)"
                                        _match = _re_sty.search(_pat, _und_text, _re_sty.DOTALL)
                                        if _match:
                                            _body = _match.group(1).strip()
                                            if _body:
                                                st.markdown(
                                                    f"<div style='background:{_sbg};border-left:4px solid {_sfg};"
                                                    f"border-radius:0 8px 8px 0;padding:10px 14px;margin:6px 0;'>"
                                                    f"<div style='font-size:0.72rem;font-weight:700;text-transform:uppercase;"
                                                    f"letter-spacing:.08em;color:{_sfg};margin-bottom:4px;'>"
                                                    f"{_sico} {_sec_name}</div>"
                                                    f"<div style='font-size:0.84rem;color:#374151;white-space:pre-wrap;'>"
                                                    f"{_body}</div></div>",
                                                    unsafe_allow_html=True
                                                )
                                else:
                                    st.info("No understanding generated yet. Click **🧠 Understand** to generate.")

                        # ── Generate Understanding Panel — Admin/Lead only ──
                        if st.session_state.get(f"sty_gen_open_{sty.id}", False) and _can_modify_style:
                            with st.expander(f"🧠 Generate Style Understanding — {sty.name}", expanded=True):
                                st.markdown(
                                    "<p style='font-size:0.82rem;color:#6b7280;'>"
                                    "The model reads ALL linked documents as a unified whole and produces a "
                                    "structured Style Intelligence Layer. Use this to validate alignment before generating.</p>",
                                    unsafe_allow_html=True
                                )
                                _doc_count = len([sd for sd in sty.style_documents if sd.document])
                                if _doc_count == 0 and not sty.custom_instructions:
                                    st.warning("⚠️ No documents or instructions linked to this style. Add content before generating understanding.")
                                else:
                                    st.caption(f"📎 {_doc_count} document(s) + {'custom instructions' if sty.custom_instructions else 'no custom instructions'} will be processed as unified context.")
                                    _sty_sys, _sty_usr, _gen_extra = render_inline_prompt_controls(
                                        db, "style",
                                        project_id=_proj_id, cluster_id=None, course_id=_crs_id,
                                        user_name=user_name, model_choice=model_choice,
                                        extra_placeholder="e.g. Pay extra attention to clinical terminology usage…",
                                        project_name=_proj_name, cluster_name=ctx.cluster_name, course_name=_crs_name,
                                    )
                                    _sty_gen_col, _sty_dl_col = st.columns([0.65, 0.35])
                                    render_prompt_download_button(
                                        db, "style",
                                        project_name=_proj_name, cluster_name=ctx.cluster_name, course_name=_crs_name,
                                        button_label="⬇️ Download Prompt Used (.md)",
                                        key=f"sty_dl_prompt_{sty.id}",
                                        use_container_width=True,
                                    )
                                    if _sty_gen_col.button("🚀 Generate Understanding", key=f"sty_gen_run_{sty.id}", type="primary", use_container_width=True):
                                        if rbac_gate(user_role, "style.understand", "Generating Style Understanding"):
                                            _mc = st.session_state.get("model_choice", "GPT-5.4")
                                            with st.spinner("Processing all documents as unified context…"):
                                                _understanding = generate_style_understanding(db, sty, _mc, _gen_extra)
                                            if _understanding.startswith("ERROR"):
                                                st.error(_understanding)
                                            else:
                                                sty.generated_summary = _understanding
                                                sty.understanding_status = "fresh"
                                                sty.updated_at = datetime.now(timezone.utc)
                                                db.commit()
                                                auto_save_instructions(
                                                    db, "style", _gen_extra,
                                                    name=f"{sty.name[:60]} validation",
                                                    project_id=_proj_id, cluster_id=None, course_id=_crs_id,
                                                    user_name=user_name,
                                                )
                                                log_event(db, "style_understanding_generated", user_name,
                                                    f"Understanding generated for style '{sty.name}'")
                                                log_audit_event(db, user_name, "style.upgraded", entity_type="style",
                                                                entity_id=sty.style_id, project_id=_proj_id, course_id=_crs_id,
                                                                metadata={"name": sty.name})
                                                st.success("✅ Style Intelligence Layer saved.")
                                                st.rerun()

                        # ── Refine / Regenerate Panel — Admin/Lead only ────
                        if st.session_state.get(f"sty_refine_open_{sty.id}", False) and _can_modify_style:
                            with st.expander(f"🔄 Refine Style Understanding — {sty.name}", expanded=True):
                                if not sty.generated_summary:
                                    st.warning("Generate an initial understanding first (click 🧠 Understand).")
                                else:
                                    st.markdown("**Current Intelligence Layer (excerpt):**")
                                    st.caption(sty.generated_summary[:400] + "…")
                                    _corr = st.text_area(
                                        "Corrections / additional guidance *",
                                        placeholder=(
                                            "e.g. The tone should be warmer, not purely academic. "
                                            "Add guidance on scenario-based openings."
                                        ),
                                        height=100, key=f"sty_refine_corr_{sty.id}"
                                    )
                                    if st.button("🔄 Regenerate Understanding", key=f"sty_regen_run_{sty.id}", type="primary"):
                                        if not _corr.strip():
                                            st.warning("Please enter corrections or guidance.")
                                        elif rbac_gate(user_role, "style.understand", "Refining Style Understanding"):
                                            _mc2 = st.session_state.get("model_choice", "GPT-5.4")
                                            with st.spinner("Refining style intelligence…"):
                                                _new_und = regenerate_style_understanding(db, sty, _mc2, _corr)
                                            if _new_und.startswith("ERROR"):
                                                st.error(_new_und)
                                            else:
                                                sty.generated_summary = _new_und
                                                sty.understanding_status = "fresh"
                                                sty.updated_at = datetime.now(timezone.utc)
                                                db.commit()
                                                log_event(db, "style_understanding_refined", user_name,
                                                    f"Understanding refined for style '{sty.name}'")
                                                st.success("✅ Refined Style Intelligence Layer saved.")
                                                st.rerun()

                        st.markdown(
                            "<hr style='border:none;border-top:1px solid #f1f5f9;margin:8px 0;'>",
                            unsafe_allow_html=True
                        )

    # =====================================================================
    # TAB 2 — DOCUMENT REGISTRY (unchanged from previous implementation)
    # =====================================================================
    with ctx_tab2:
        # ── How-to Banner ────────────────────────────────────────────────
        with st.expander("💡 How to reference documents in prompts", expanded=False):
            st.markdown(
                r"""
**Document-Aware Generation** — mention any uploaded filename in backticks inside your topic or prompt.

**Examples:**
- `` `Instructional_Design_Spec_v2.pdf` `` — Generate a lesson summary based on this file
- `` `Course_Outline.docx` `` — Use this as reference material for the module

The system will detect the filename, retrieve the content, and inject it as structured context.
                """
            )

        # ── Metrics Bar — count queries only, no full-table load ─────────
        _doc_total    = document_repository.count_documents(db)
        _doc_active   = document_repository.count_active_documents(db)
        _doc_archived = document_repository.count_archived_documents(db)
        _doc_tags     = document_repository.list_distinct_document_tags(db)

        _m1, _m2, _m3, _m4 = st.columns(4)
        _m1.metric("📄 Total",     _doc_total)
        _m2.metric("✅ Active",    _doc_active)
        _m3.metric("🗄️ Archived", _doc_archived)
        _m4.metric("🏷️ Types",    len(_doc_tags))

        st.divider()

        # ── Upload Panel ─────────────────────────────────────────────────
        with st.expander("⬆️ Upload New Document", expanded=False):
            with st.form(key="ctx_upload_form_v2"):
                _uf1, _uf2 = st.columns([0.6, 0.4])
                with _uf1:
                    ctx_new_file = st.file_uploader(
                        "Select File",
                        type=["pdf", "docx", "xlsx", "txt"],
                        key="ctx_new_file_uploader",
                    )
                with _uf2:
                    ctx_doc_type = st.selectbox(
                        "Document Type",
                        ["guidelines", "checklist", "chapter", "specification", "outline",
                         "assessment", "rubric", "reference", "style_reference", "general"],
                    )
                _ub1, _ub2, _ub3 = st.columns([0.25, 0.5, 0.25])
                ctx_commit_btn = _ub2.form_submit_button("📁 Add to Database", use_container_width=True, type="primary")

            if ctx_commit_btn:
                if not ctx_new_file:
                    st.warning("Please select a file.")
                else:
                    _ex = document_repository.get_document_by_filename(db, ctx_new_file.name)
                    if _ex and (_ex.status or "active") == "active":
                        st.warning(f"'{ctx_new_file.name}' already exists. Use Update action below.")
                    else:
                        ctx_content = ""
                        _fn = ctx_new_file.name.lower()
                        try:
                            if _fn.endswith(".pdf"):
                                reader = PdfReader(ctx_new_file)
                                ctx_content = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
                            elif _fn.endswith(".docx"):
                                _doc = DocxDocument(ctx_new_file)
                                ctx_content = "\n".join([p.text for p in _doc.paragraphs])
                            elif _fn.endswith(".xlsx"):
                                _df = pd.read_excel(ctx_new_file)
                                ctx_content = _df.to_csv(index=False)
                            else:
                                ctx_content = ctx_new_file.read().decode("utf-8", errors="ignore")
                        except Exception as _e:
                            st.error(f"❌ {_e}")
                        if ctx_content:
                            _now = datetime.now(timezone.utc)
                            if _ex:
                                _ex.content = ctx_content; _ex.file_type = ctx_new_file.type
                                _ex.doc_tag = ctx_doc_type; _ex.uploaded_by = user_name
                                _ex.updated_at = _now; _ex.status = "active"; db.commit()
                                st.success(f"✅ '{ctx_new_file.name}' reactivated.")
                            else:
                                db.add(Document(
                                    filename=ctx_new_file.name, file_type=ctx_new_file.type,
                                    doc_tag=ctx_doc_type, content=ctx_content,
                                    uploaded_by=user_name, uploaded_at=_now,
                                    updated_at=_now, status="active"
                                )); db.commit()
                                st.success(f"✅ '{ctx_new_file.name}' added ({len(ctx_content):,} chars).")
                            log_event(db, "document_upload", user_name,
                                f"Uploaded '{ctx_new_file.name}' as {ctx_doc_type}",
                                {"filename": ctx_new_file.name, "chars": len(ctx_content)})
                            st.rerun()

        st.divider()
        st.markdown("#### 🗂️ Document Registry")

        # Filters — DB-side, resets page on change
        _cf1, _cf2, _cf3 = st.columns([0.35, 0.35, 0.3])
        _ctx_filter_type   = _cf1.selectbox("Filter by Type",   ["All"] + _doc_tags,                   key="ctx_filter_type")
        _ctx_filter_status = _cf2.selectbox("Filter by Status", ["All", "active", "archived"],          key="ctx_filter_status")
        _ctx_search        = _cf3.text_input("🔍 Search filename", placeholder="e.g. Spec_v2",          key="ctx_search_q")

        _flt_status = None if _ctx_filter_status == "All" else _ctx_filter_status
        _flt_tag    = None if _ctx_filter_type   == "All" else _ctx_filter_type
        _flt_search = _ctx_search.strip() or None

        _doc_filtered_total = document_repository.count_documents_filtered(
            db, status=_flt_status, tag=_flt_tag, search=_flt_search,
        )

        if _doc_filtered_total == 0:
            st.info("No documents match your filters.")
        else:
            _doc_offset, _ = _paginate_docs(_doc_filtered_total, DOCUMENT_PAGE_SIZE, "doc_registry_page")
            filtered_ctx_docs = document_repository.list_documents_filtered(
                db,
                status=_flt_status,
                tag=_flt_tag,
                search=_flt_search,
                limit=DOCUMENT_PAGE_SIZE,
                offset=_doc_offset,
            )
            for _doc in filtered_ctx_docs:
                _doc_status   = _doc.status or "active"
                _status_color = "#10b981" if _doc_status == "active" else "#6b7280"
                _status_bg    = "#ecfdf5" if _doc_status == "active" else "#f3f4f6"
                _updated_str  = _doc.updated_at.strftime("%Y-%m-%d %H:%M") if _doc.updated_at else "—"
                _uploaded_str = _doc.uploaded_at.strftime("%Y-%m-%d %H:%M") if _doc.uploaded_at else "—"
                _chars        = len(_doc.content) if _doc.content else 0

                with st.container():
                    _rc1, _rc2, _rc3, _rc4, _rc5, _rc6 = st.columns([0.28, 0.12, 0.14, 0.14, 0.1, 0.22])
                    _rc1.markdown(
                        f"<div style='font-weight:600;font-size:0.875rem;color:#111827;padding-top:6px;'>"
                        f"📄 {_doc.filename}</div>"
                        f"<div style='font-size:0.72rem;color:#9ca3af;'>#{_doc.id} · {_chars:,} chars</div>",
                        unsafe_allow_html=True
                    )
                    _rc2.markdown(
                        f"<div style='padding-top:8px;'><span style='background:#eef2ff;color:#4338ca;"
                        f"border:1px solid #c7d2fe;padding:2px 8px;border-radius:12px;"
                        f"font-size:0.72rem;font-weight:600;text-transform:uppercase;'>"
                        f"{_doc.doc_tag or 'general'}</span></div>", unsafe_allow_html=True
                    )
                    _rc3.markdown(
                        f"<div style='font-size:0.78rem;color:#6b7280;padding-top:8px;'>"
                        f"📅 <strong>Uploaded</strong><br>{_uploaded_str}</div>", unsafe_allow_html=True
                    )
                    _rc4.markdown(
                        f"<div style='font-size:0.78rem;color:#6b7280;padding-top:8px;'>"
                        f"🔄 <strong>Updated</strong><br>{_updated_str}</div>", unsafe_allow_html=True
                    )
                    _rc5.markdown(
                        f"<div style='padding-top:8px;'><span style='background:{_status_bg};"
                        f"color:{_status_color};border:1.5px solid {_status_color}44;"
                        f"padding:2px 8px;border-radius:12px;font-size:0.72rem;font-weight:700;"
                        f"text-transform:uppercase;'>{'✅ active' if _doc_status == 'active' else '🗄 archived'}"
                        f"</span></div>", unsafe_allow_html=True
                    )
                    with _rc6:
                        _ac = st.columns(3)
                        if _ac[0].button("👁️", key=f"ctx_view_{_doc.id}", help="View content"):
                            st.session_state[f"ctx_view_open_{_doc.id}"] = not st.session_state.get(f"ctx_view_open_{_doc.id}", False)
                        if _ac[1].button("✏️", key=f"ctx_update_{_doc.id}", help="Update"):
                            st.session_state[f"ctx_update_open_{_doc.id}"] = not st.session_state.get(f"ctx_update_open_{_doc.id}", False)
                        if _doc_status == "active":
                            if _ac[2].button("🗄️", key=f"ctx_archive_{_doc.id}", help="Archive"):
                                _doc.status = "archived"; _doc.updated_at = datetime.now(timezone.utc)
                                db.commit(); st.toast(f"🗄️ '{_doc.filename}' archived."); st.rerun()
                        else:
                            if _ac[2].button("🗑️", key=f"ctx_delete_{_doc.id}", help="Delete"):
                                db.delete(_doc); db.commit(); st.toast(f"🗑️ Deleted."); st.rerun()

                if st.session_state.get(f"ctx_view_open_{_doc.id}", False):
                    with st.expander(f"📖 Preview — {_doc.filename}", expanded=True):
                        st.text_area("Content (first 3,000 chars)", value=(_doc.content or "")[:3000],
                                     height=200, disabled=True, key=f"ctx_preview_{_doc.id}")
                        st.markdown(
                            f"<div style='background:#f0f9ff;border:1px solid #bae6fd;border-radius:8px;"
                            f"padding:8px 12px;font-size:0.82rem;color:#0369a1;margin-top:6px;'>"
                            f"💡 Reference in prompt: <code>`{_doc.filename}`</code></div>",
                            unsafe_allow_html=True
                        )

                if st.session_state.get(f"ctx_update_open_{_doc.id}", False):
                    with st.expander(f"✏️ Update — {_doc.filename}", expanded=True):
                        with st.form(key=f"ctx_update_form_{_doc.id}"):
                            _upd_type = st.selectbox("Document Type",
                                ["guidelines","checklist","chapter","specification","outline",
                                 "assessment","rubric","reference","style_reference","general"],
                                index=["guidelines","checklist","chapter","specification","outline",
                                       "assessment","rubric","reference","style_reference","general"].index(
                                       _doc.doc_tag if _doc.doc_tag in
                                       ["guidelines","checklist","chapter","specification","outline",
                                        "assessment","rubric","reference","style_reference","general"] else "general"),
                                key=f"ctx_upd_type_{_doc.id}"
                            )
                            _upd_file = st.file_uploader("Replacement file (optional)",
                                type=["pdf","docx","xlsx","txt"], key=f"ctx_upd_file_{_doc.id}")
                            if st.form_submit_button("💾 Save", type="primary"):
                                _now_upd = datetime.now(timezone.utc)
                                _doc.doc_tag = _upd_type; _doc.updated_at = _now_upd; _doc.status = "active"
                                if _upd_file:
                                    _uc = ""
                                    _ufn2 = _upd_file.name.lower()
                                    try:
                                        if _ufn2.endswith(".pdf"):
                                            _r2 = PdfReader(_upd_file)
                                            _uc = "\n".join([_p2.extract_text() for _p2 in _r2.pages if _p2.extract_text()])
                                        elif _ufn2.endswith(".docx"):
                                            _ud2 = DocxDocument(_upd_file); _uc = "\n".join([_p2.text for _p2 in _ud2.paragraphs])
                                        elif _ufn2.endswith(".xlsx"):
                                            _udf2 = pd.read_excel(_upd_file); _uc = _udf2.to_csv(index=False)
                                        else:
                                            _uc = _upd_file.read().decode("utf-8", errors="ignore")
                                    except Exception as _ue2:
                                        st.error(f"❌ {_ue2}")
                                    if _uc:
                                        _doc.content = _uc; _doc.file_type = _upd_file.type
                                        _doc.filename = _upd_file.name; _doc.uploaded_by = user_name
                                db.commit()
                                st.success(f"✅ Updated."); st.session_state[f"ctx_update_open_{_doc.id}"] = False; st.rerun()

                st.markdown("<hr style='border:none;border-top:1px solid #f1f5f9;margin:4px 0;'>", unsafe_allow_html=True)

        if _doc_active > 0:
            st.divider()
            st.markdown("#### 📋 Quick Reference Card")
            _ref_docs = document_repository.list_active_documents(db, limit=100)
            _ref_lines = [f"• `{d.filename}` — {d.doc_tag or 'general'}" for d in _ref_docs]
            if _doc_active > 100:
                _ref_lines.append(f"… and {_doc_active - 100} more (use filters above to find them)")
            st.code("\n".join(_ref_lines), language=None)

    # ── Prompt Management ─────────────────────────────────────────────────
    st.divider()
    st.markdown(
        "<div style='font-size:0.75rem;font-weight:700;text-transform:uppercase;"
        "letter-spacing:.1em;color:#6366f1;margin-bottom:4px;'>📚 Prompt Library</div>"
        "<p style='font-size:0.82rem;color:#6b7280;margin-bottom:8px;'>"
        "Create, edit, version, and improve prompt assets tagged <code>style</code>. "
        "These are available across all generation components.</p>",
        unsafe_allow_html=True,
    )
    render_prompt_panel(
        db,
        component="style",
        user_name=user_name,
        model_choice=model_choice,
        default_system=(
            "You are an expert instructional designer with deep expertise in instructional style "
            "and tone. Analyse the provided style guidelines and documents."
        ),
        default_user=(
            "Apply the following style guidelines to the content generation task:\n\n{style_context}"
        ),
    )

