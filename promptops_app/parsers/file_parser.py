"""File parser — extract plain text from uploaded PDF / DOCX / XLSX / TXT files.

Contains no Streamlit calls — safe to use from ThreadPoolExecutor workers
and from future FastAPI endpoints.

Extracted from core/shared.py (Phase 3 refactoring).
"""

import io
import logging
import time

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

_log = logging.getLogger(__name__)


def parse_uploaded_file(uploaded_file) -> tuple:
    """Extract text from a Streamlit UploadedFile (or any file-like object).

    Returns:
        (filename: str, content: str, error_msg: str | None)

    Safe to call from a ThreadPoolExecutor worker — no st.* calls.
    Supports: PDF, DOCX, PPTX, XLSX, TXT (and plain text variants).
    """
    name = uploaded_file.name
    fn   = name.lower()
    start = time.monotonic()

    try:
        if fn.endswith(".pdf"):
            reader  = PdfReader(uploaded_file)
            content = "\n".join(
                page.extract_text()
                for page in reader.pages
                if page.extract_text()
            )
        elif fn.endswith(".docx"):
            doc     = DocxDocument(uploaded_file)
            content = "\n".join(p.text for p in doc.paragraphs)
        elif fn.endswith(".doc"):
            return name, "", (
                "Legacy .doc format is not supported. "
                "Please save the file as .docx and re-upload."
            )
        elif fn.endswith(".pptx"):
            try:
                from pptx import Presentation as _Pptx  # python-pptx
                _prs = _Pptx(io.BytesIO(uploaded_file.read()))
                content = "\n".join(
                    shape.text
                    for slide in _prs.slides
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                )
            except ImportError:
                return name, "", (
                    "python-pptx is not installed. "
                    "Ask your admin to add python-pptx to requirements.txt."
                )
        elif fn.endswith(".ppt"):
            return name, "", (
                "Legacy .ppt format is not supported. "
                "Please save the file as .pptx and re-upload."
            )
        elif fn.endswith(".xlsx"):
            content = pd.read_excel(uploaded_file).to_csv(index=False)
        else:
            content = uploaded_file.read().decode("utf-8", errors="ignore")

        elapsed = time.monotonic() - start
        _log.info(
            "file_parser.parse  name=%r  ext=%r  chars=%d  duration=%.3fs",
            name, fn.rsplit(".", 1)[-1] if "." in fn else "txt",
            len(content), elapsed,
        )
        return name, content, None

    except Exception as exc:
        elapsed = time.monotonic() - start
        _log.warning(
            "file_parser.parse FAILED  name=%r  duration=%.3fs  error=%s",
            name, elapsed, exc,
        )
        return name, "", str(exc)


# Backward-compatible private alias used throughout the codebase
_parse_uploaded_file = parse_uploaded_file
